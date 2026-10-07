"""Generate editable XLSX, recalculate with bundled LibreOffice, and publish reproducible JSON.

Run with the bundled Python documented in README.md. Explicitly illustrative inputs remain labelled.
No private analytics, measured conversion or approved quote is fabricated.
"""
from pathlib import Path
import argparse, json, shutil, subprocess, tempfile
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

HERE = Path(__file__).resolve().parent
SOFFICE = Path('/Users/jreye/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
DELIVERY = Path('/Users/jreye/Documents/Banco Pichincha Peru — Prospect Case/02-Uso-interno/modelo')
NAME = 'Pibank-caso-negocio-2027.xlsx'
NUM = '#,##0.0;(#,##0.0);"—"'
USD = '"US$"#,##0;("US$"#,##0);"—"'
PCT = '0.0%;(0.0%);"—"'
RATE = '0.00%;(0.00%);"—"'

def monthdate(start, offset):
    n = start.year * 12 + start.month - 1 + offset
    return date(n // 12, n % 12 + 1, 1)

def validate_inputs(cfg):
    assert cfg['months'] == 6
    assert cfg['selected_case'] in (1, 2, 3)
    assert cfg['demand_monthly'] >= 0 and cfg['ticket_usd'] >= 0
    for cs in cfg['cases']:
        assert len(cs['extra_capture']) == len(cs['reference_capture']) == 6
        for key in ('opening_rate','funding_rate','retention30','retention60','retention90','external_share'):
            assert 0 <= cs[key] <= 1
        assert cs['retention90'] <= cs['retention60'] <= cs['retention30']
        for r, p in zip(cs['reference_capture'], cs['extra_capture']):
            assert 0 <= r <= 1 and 0 <= p <= 1 and r+p <= 1

def simulate(cfg, case_index):
    cs=cfg['cases'][case_index-1]; demand=cfg['demand_monthly']*cfg['relevance_eligibility_factor']; ticket=cfg['ticket_usd']
    visits=[demand*x for x in cs['extra_capture']]
    openings=[v*cs['opening_rate'] for v in visits]
    funded=[o*cs['funding_rate'] for o in openings]
    cohort=[]
    for i,n in enumerate(funded):
        balances=[]
        for j in range(6):
            age=j-i
            retention=0 if age<0 else 1 if age==0 else cs['retention30'] if age==1 else cs['retention60'] if age==2 else cs['retention90']
            balances.append(n*ticket*retention)
        cohort.append(balances)
    stock=[sum(cohort[i][j] for i in range(6)) for j in range(6)]
    costs=cfg['monthly_cost_usd']
    ready=cfg['finance_validated'] and cfg['costs_validated'] and cfg['net_annual_value_per_usd_external_balance'] is not None and all(isinstance(c,(int,float)) for c in costs) and sum(costs)>0
    value=[s*cs['external_share']*cfg['net_annual_value_per_usd_external_balance']/12 for s in stock] if ready else None
    payback=next((i+1 for i in range(6) if sum(value[:i+1])>=sum(costs[:i+1])), None) if ready else None
    return {'case':cs['name'],'nature':'Ilustración condicionada, no forecast ni resultado medido',
            'visits_incremental':sum(visits),'openings_incremental':sum(openings),'funded_incremental':sum(funded),
            'first_deposit_usd_equivalent':sum(funded)*ticket,'balance_end_usd_equivalent':stock[-1],
            'external_balance_end_usd_equivalent':stock[-1]*cs['external_share'],
            'account_goal_share':sum(funded)/cfg['goal_accounts'],
            'roi':(sum(value)-sum(costs))/sum(costs) if ready else None,'payback_month':payback,
            'finance_status':'N/D hasta valoración y costes validados' if not ready else 'Validado como input de prueba',
            'monthly':[{'month':monthdate(date.fromisoformat(cfg['start_month']),i).isoformat(),'visits_incremental':visits[i],
                        'openings_incremental':openings[i],'funded_incremental':funded[i],
                        'balance_end_usd_equivalent':stock[i],'external_balance_end_usd_equivalent':stock[i]*cs['external_share']} for i in range(6)],
            'cohort_balances':cohort}

def value(ws, address, data, fmt=None, input_cell=False):
    c=ws[address]; c.value=data
    color='1A1A1A'
    if input_cell:color='0000CC'
    elif isinstance(data,str) and data.startswith('=') and '!' in data:color='008050'
    c.font=Font(name='Arial',size=10,color=color)
    c.alignment=Alignment(vertical='center',wrap_text=False)
    if fmt:c.number_format=fmt
    if input_cell:c.fill=PatternFill('solid',fgColor='FFF2CC')
    return c

def textline(ws,row,text):
    ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=11)
    c=value(ws,f'C{row}',text);c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[row].height=30

def setup(wb,name,title,end_row=35):
    ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
    for c in 'AB':ws.column_dimensions[c].width=2
    ws.column_dimensions['C'].width=40
    ws.column_dimensions['D'].width=18;ws.column_dimensions['E'].width=3
    for c in 'FGHIJK':ws.column_dimensions[c].width=15
    ws.row_dimensions[2].height=25
    value(ws,'C2',title).font=Font(name='Arial',size=14,bold=True,color='18294A')
    value(ws,'C4','Escenario seleccionado')
    value(ws,'D4',"='Supuestos'!D5")
    ws['D4'].border=Border(bottom=Side(style='dotted',color='7E8A9E'))
    ws['D4'].alignment=Alignment(horizontal='center',vertical='center')
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True
    ws.print_area=f'C1:K{end_row}'
    ws.page_margins.left=ws.page_margins.right=0.25
    ws.page_margins.top=ws.page_margins.bottom=0.3
    for r in range(5,end_row+1):ws.row_dimensions[r].height=20
    return ws

def header(ws,row,labels,start_col=3):
    for col,label in enumerate(labels,start_col):
        c=value(ws,f'{openpyxl.utils.get_column_letter(col)}{row}',label)
        c.fill=PatternFill('solid',fgColor='20385F');c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[row].height=30

def recalculate(source,dest):
    dest.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='pibank-lo-profile-') as profile:
        command=[str(SOFFICE),f'-env:UserInstallation={Path(profile).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(dest),str(source)]
        p=subprocess.run(command,capture_output=True,text=True,timeout=90)
        if p.returncode or not (dest/source.name).exists():raise RuntimeError(f'Recalculation failed: {p.returncode}: {p.stdout} {p.stderr}')
    return dest/source.name

def make_book(cfg, output):
    wb=openpyxl.Workbook();wb.remove(wb.active)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    s=setup(wb,'Resumen','Pibank: caso de captación 2027',38)
    a=setup(wb,'Supuestos','Inputs editables y escenarios',73)
    m=setup(wb,'Mensual','Captación incremental: seis meses',37)
    co=setup(wb,'Cohortes','Saldo de cohortes incrementales',34)
    se=setup(wb,'Sensibilidad','Sensibilidad y escala requerida',55)
    f=setup(wb,'Fuentes','Procedencia y límites de inputs',36)
    demand=setup(wb,'Demanda','Estimaciones de mercado PE/es: fuente histórica',30)
    s.sheet_properties.tabColor='20385F';a.sheet_properties.tabColor='687DA0'
    for ws in (a,m,co):ws.freeze_panes='F8'
    # One case selector and global controls.
    value(a,'C4','Selector escenario (1 / 2 / 3)');value(a,'D4',cfg['selected_case'],input_cell=True)
    value(a,'C5','Nombre activo');value(a,'D5','=CHOOSE(D4,"Conservador","Base","Favorable")')
    textline(a,6,'1 Conservador, 2 Base, 3 Favorable. Los nombres distinguen pruebas ilustrativas, no probabilidades.')
    dv=DataValidation(type='whole',operator='between',formula1=1,formula2=3);dv.error='Usa 1, 2 o 3';dv.showErrorMessage=True;a.add_data_validation(dv);dv.add(a['D4'])
    globals_=[(8,'Meta cuentas provisional',cfg['goal_accounts'],NUM),(9,'Meta USD provisional',cfg['goal_usd'],USD),(10,'Ticket USD equivalente provisional',cfg['ticket_usd'],USD),
              (11,'Inicio ilustrativo',date.fromisoformat(cfg['start_month']),'mmm-yy'),(13,'Factor relevancia/eligibilidad ilustrativo',cfg['relevance_eligibility_factor'],PCT),(14,'Finance validó contribución neta','Sí' if cfg['finance_validated'] else 'No',None),
              (15,'Contribución neta anual / USD externo',cfg['net_annual_value_per_usd_external_balance'],RATE),(16,'Costos validados','Sí' if cfg['costs_validated'] else 'No',None)]
    for row,label,v,fmt in globals_:value(a,f'C{row}',label);value(a,f'D{row}',v,fmt,input_cell=True)
    textline(a,12,'Meta USD inferida de 15.200 × US$2.500. Moneda y periodo de la meta no confirmados.')
    for row in (14,16):
        yesno=DataValidation(type='list',formula1='"Sí,No"');a.add_data_validation(yesno);yesno.add(a[f'D{row}'])
    start=date.fromisoformat(cfg['start_month']);cols=list('FGHIJK')
    for i,col in enumerate(cols):value(a,f'{col}18',f'=EDATE($D$11,{i})','mmm-yy')
    value(a,'C17','Capacidad con reserva (horas)');value(a,'D17',cfg.get('delivery_capacity_hours_with_reserve'),NUM)
    value(a,'C19','Inversión cliente USD / mes (pendiente)')
    for i,col in enumerate(cols):value(a,f'{col}19',cfg['monthly_cost_usd'][i],USD,input_cell=True)
    value(a,'C20','Inversión seis meses, si completa');value(a,'D20','=IF(COUNT(F19:K19)=6,SUM(F19:K19),"N/D")',USD)
    value(a,'C21','Costo interno Efeonce USD (pendiente)');value(a,'D21',cfg.get('efeonce_internal_cost_usd'),USD,input_cell=True)
    driver_rows={'extra_capture':23,'reference_capture':28,'opening_rate':33,'funding_rate':38,'demand':43,'retention30':48,'retention60':53,'retention90':58,'external_share':63}
    labels={'extra_capture':'Captura adicional de demanda','reference_capture':'Captura referencia sin programa','opening_rate':'Visita a apertura completada','funding_rate':'Apertura a primer fondeo','demand':'Demanda mensual (búsquedas, no personas)','retention30':'Saldo relativo a 30 días','retention60':'Saldo relativo a 60 días','retention90':'Saldo relativo a 90 días y después','external_share':'Proporción externa del saldo incremental'}
    for key,row in driver_rows.items():
        value(a,f'C{row}',labels[key]+' activo')
        for i,col in enumerate(cols):value(a,f'{col}{row}',f'=IF(CHOOSE($D$4,COUNT({col}{row+1}),COUNT({col}{row+2}),COUNT({col}{row+3}))=1,CHOOSE($D$4,{col}{row+1},{col}{row+2},{col}{row+3}),"N/D")',NUM if key=='demand' else RATE)
        for case_index,cs in enumerate(cfg['cases'],1):
            value(a,f'C{row+case_index}',cs['name'])
            raw=cfg['demand_monthly'] if key=='demand' else cs[key]
            for i,col in enumerate(cols):value(a,f'{col}{row+case_index}',raw[i] if isinstance(raw,list) else raw,NUM if key=='demand' else RATE,input_cell=True)
    textline(a,68,'Amarillo/azul: editable. Verde: vínculo. Negro: cálculo. Demanda proxy histórica; conversión y captura ilustrativas.')
    textline(a,69,'Retención es proporción de saldo por cuenta original, incluye salida de fondos; no multiplicar por abandono otra vez.')
    textline(a,70,'Valor anual neto: UN método Finance que ya descuenta costos del fondeo. No sumar FTP y ahorro duplicados.')
    textline(a,71,'USD representa equivalencia de la meta. La cuenta pública vigente es Soles; no se modela tipo de cambio.')
    textline(a,72,cfg['start_status'])
    textline(a,73,cfg.get('client_investment_status','Inversión pendiente. No asumir cero ni tarifa aprobada.'))
    # Monthly build consumes only active assumptions.
    header(m,7,['Indicador','Unidad','','Mes 1','Mes 2','Mes 3','Mes 4','Mes 5','Mes 6'])
    for col in cols:value(m,f'{col}8',f"='Supuestos'!{col}18",'mmm-yy')
    monthly_labels={9:('Proxy demanda histórica trasladada','búsquedas'),10:('Captura referencia','%'),11:('Captura adicional programa','%'),12:('Visitas referencia','visitas'),13:('Visitas adicionales','visitas'),14:('Visitas con programa','visitas'),15:('Demanda × relevancia ilustrativa','búsquedas'),16:('Visita → apertura','%'),17:('Aperturas referencia','cuentas'),18:('Aperturas con programa','cuentas'),19:('Aperturas incrementales','cuentas'),20:('Apertura → fondeo','%'),21:('Fondeadas referencia','cuentas'),22:('Fondeadas con programa','cuentas'),23:('Fondeadas incrementales','cuentas'),25:('Saldo incremental al cierre','USD eq.'),26:('Saldo externo incremental','USD eq.'),27:('Inversión cliente pendiente','USD'),28:('Inversión acumulada','USD'),29:('Contribución neta mensual','USD'),30:('Contribución acumulada','USD'),31:('ROI acumulado','%'),32:('Valor menos inversión acumulado','USD'),34:('Recuperación alcanzada','0/1')}
    for row,(label,unit) in monthly_labels.items():value(m,f'C{row}',label);value(m,f'D{row}',unit)
    for j,col in enumerate(cols):
        refs={9:43,10:28,11:23,16:33,20:38}
        for row,src in refs.items():value(m,f'{col}{row}',f"='Supuestos'!{col}{src}",NUM if row==9 else RATE)
        forms={15:f"{col}9*'Supuestos'!$D$13",12:f'{col}15*{col}10',13:f'{col}15*{col}11',14:f'SUM({col}12:{col}13)',17:f'{col}12*{col}16',18:f'{col}14*{col}16',19:f'{col}18-{col}17',21:f'{col}17*{col}20',22:f'{col}18*{col}20',23:f'{col}22-{col}21',25:f"'Cohortes'!{openpyxl.utils.get_column_letter(11+j)}15",26:f'{col}25*\'Supuestos\'!{col}63',
               27:f'IF(COUNT(\'Supuestos\'!{col}19)=1,\'Supuestos\'!{col}19,"N/D")',28:f'IF(COUNT($F27:{col}27)={j+1},SUM($F27:{col}27),"N/D")',
               29:f'IF(AND(\'Supuestos\'!$D$14="Sí",COUNT(\'Supuestos\'!$D$15,{col}26)=2),{col}26*\'Supuestos\'!$D$15/12,"N/D")',
               30:f'IF(COUNT($F29:{col}29)={j+1},SUM($F29:{col}29),"N/D")',
               31:f'IF(AND(\'Supuestos\'!$D$14="Sí",\'Supuestos\'!$D$16="Sí",COUNT({col}28,{col}30)=2,{col}28>0),({col}30-{col}28)/{col}28,"N/D")',
               32:f'IF(AND(\'Supuestos\'!$D$14="Sí",\'Supuestos\'!$D$16="Sí",COUNT({col}28,{col}30)=2,{col}28>0),{col}30-{col}28,"N/D")',34:f'IF(COUNT({col}32)=1,IF({col}32>=0,1,0),"N/D")'}
        guards={12:f'{col}15,{col}10',13:f'{col}15,{col}11',14:f'{col}12,{col}13',15:f"{col}9,'Supuestos'!$D$13",17:f'{col}12,{col}16',18:f'{col}14,{col}16',19:f'{col}18,{col}17',21:f'{col}17,{col}20',22:f'{col}18,{col}20',23:f'{col}22,{col}21',26:f"{col}25,'Supuestos'!{col}63"}
        for row,formula in forms.items():
            if row in guards:formula=f'IF(COUNT({guards[row]})=2,{formula},"N/D")'
            value(m,f'{col}{row}','='+formula,PCT if row==31 else USD if row in (25,26,27,28,29,30,32) else NUM)
    textline(m,36,'Referencia también es ilustrativa. No representa tráfico actual ni pérdidas históricas desde el lanzamiento.')
    textline(m,37,'Contribución valorada sólo sobre saldo externo. Ninguna apertura se atribuye a una cita de IA.')
    # Cohort balances: stock not sum of stocks. Age zero is entry deposit, 30d applies following month.
    co.column_dimensions['C'].width=16;co.column_dimensions['D'].width=16;co.column_dimensions['E'].width=16
    for col in 'FGHIJKLMNOP':co.column_dimensions[col].width=14
    co.print_area='C1:P34'
    header(co,7,['Cohorte','Fondeadas inc.','Ticket USD eq.','Relativo 30d','Relativo 60d','Relativo 90d','','','M1','M2','M3','M4','M5','M6'])
    for i,col in enumerate(cols):
        row=8+i;value(co,f'C{row}',f"='Supuestos'!{col}18",'mmm-yy');value(co,f'D{row}',f"='Mensual'!{col}23",NUM);value(co,f'E{row}',"=IF(COUNT('Supuestos'!$D$10)=1,'Supuestos'!$D$10,\"N/D\")",USD)
        for cc,sr in [('F',48),('G',53),('H',58)]:value(co,f'{cc}{row}',f"='Supuestos'!{col}{sr}",PCT)
        for j in range(6):
            cc=openpyxl.utils.get_column_letter(11+j);age=j-i
            ret="F" if age==1 else "G" if age==2 else "H"
            formula='=0' if age<0 else f'=IF(COUNT($D{row},$E{row})=2,$D{row}*$E{row},"N/D")' if age==0 else f'=IF(COUNT($D{row},$E{row},${ret}{row})=3,$D{row}*$E{row}*${ret}{row},"N/D")'
            value(co,f'{cc}{row}',formula,USD)
    value(co,'C15','Saldo cierre USD eq.')
    for j in range(6):cc=openpyxl.utils.get_column_letter(11+j);value(co,f'{cc}15',f'=IF(COUNT({cc}8:{cc}13)=6,SUM({cc}8:{cc}13),"N/D")',USD)
    textline(co,18,'Saldo relativo incluye ceros/retiros por cuenta original. Curva ilustrativa; sin datos bancarios observados.')
    header(co,21,['Cohorte','Cuentas iniciales','Saldo a 30 días','Saldo a 60 días','Saldo a 90 días'])
    for i in range(6):
        row=22+i;src=8+i;value(co,f'C{row}',f'=C{src}','mmm-yy');value(co,f'D{row}',f'=D{src}',NUM)
        for col,age,ret in [('E',1,'F'),('F',2,'G'),('G',3,'H')]:value(co,f'{col}{row}',f'=IF(COUNT($D{src},$E{src},${ret}{src})=3,$D{src}*$E{src}*${ret}{src},"N/D")' if i+age<6 else 'Fuera 6m',USD)
    textline(co,29,'30/60/90 días son proxies mensuales desde primer fondeo. Los que caen después del semestre figuran Fuera 6m.')
    textline(co,30,'La suma de depósitos iniciales es entrada bruta modelada. El saldo cierre es stock retenido; no se suman stocks.')
    textline(co,31,'Cuentas fraccionarias son expectativas matemáticas. No representan personas reales ni cuentas observadas.')
    # Simple self-contained sensitivities, no copied complex model or inert formulas.
    header(se,7,['Banda meta','Cuentas requeridas','Visitas necesarias','Saldo entrada USD eq.'])
    for row,share in [(8,.01),(9,.05),(10,.10)]:
        value(se,f'C{row}',share,PCT,input_cell=True);value(se,f'D{row}',f"=C{row}*'Supuestos'!$D$8",NUM)
        value(se,f'E{row}',f'=IF(COUNT(\'Supuestos\'!$F$33,\'Supuestos\'!$F$38)=2,IF(\'Supuestos\'!$F$33*\'Supuestos\'!$F$38>0,D{row}/(\'Supuestos\'!$F$33*\'Supuestos\'!$F$38),"N/D"),"N/D")',NUM)
        value(se,f'F{row}',f"=D{row}*'Supuestos'!$D$10",USD)
    se.column_dimensions['E'].width=20
    textline(se,12,'Bandas de escala: no son compromiso de contribución. Tasas de M1 se usan sólo en este cálculo top-down.')
    textline(se,14,'Sensibilidad: cuentas incrementales con las visitas adicionales del caso activo.')
    header(se,16,['Apertura / Fondeo','','',.4,.5,.65,.75,.85])
    for row,rate in enumerate([.01,.015,.02,.025,.03],17):
        value(se,f'C{row}',rate,PCT,input_cell=True)
        for col in 'FGHIJ':
            se[f'{col}16'].number_format=PCT
            value(se,f'{col}{row}',f"=IF(COUNT('Mensual'!$F$13:$K$13)=6,SUM('Mensual'!$F$13:$K$13)*$C{row}*{col}$16,\"N/D\")",NUM)
    header(se,24,['Retraso meses','Visitas dentro 6m','Cuentas fondeadas'])
    for row,delay in [(25,0),(26,1),(27,2)]:
        value(se,f'C{row}',delay,input_cell=True)
        value(se,f'D{row}',f'=IF(AND(C{row}>=0,C{row}<6,COUNT(\'Mensual\'!$F$13:INDEX(\'Mensual\'!$F$13:$K$13,1,6-C{row}))=6-C{row}),SUM(\'Mensual\'!$F$13:INDEX(\'Mensual\'!$F$13:$K$13,1,6-C{row})),"N/D")',NUM)
        value(se,f'E{row}',f'=IF(COUNT(D{row},\'Supuestos\'!$F$33,\'Supuestos\'!$F$38)=3,D{row}*\'Supuestos\'!$F$33*\'Supuestos\'!$F$38,"N/D")',NUM)
    textline(se,29,'Retraso desplaza la curva y recorta los últimos meses dentro del horizonte. Conversiones M1 constantes en esta prueba.')
    header(se,32,['Origen externo','Saldo externo cierre USD eq.'])
    for row,share in [(33,.4),(34,.7),(35,1)]:value(se,f'C{row}',share,PCT,input_cell=True);value(se,f'D{row}',f"=IF(COUNT('Mensual'!$K$25)=1,'Mensual'!$K$25*C{row},\"N/D\")",USD)
    textline(se,37,'Retención y ticket se editan en Supuestos y actualizan Cohortes/Resumen. No equivalen a margen o revenue.')
    header(se,40,['Inversión hipotética USD','Valor por cuenta para equilibrio','Naturaleza'])
    for row,cost in enumerate(cfg.get('investment_sensitivity_usd',[10000,25000,50000]),41):
        value(se,f'C{row}',cost,USD,input_cell=True)
        value(se,f'D{row}',f'=IF(AND(COUNT(\'Resumen\'!$D$10)=1,\'Resumen\'!$D$10>0),C{row}/\'Resumen\'!$D$10,"N/D")',USD)
        value(se,f'E{row}','Sensibilidad; no fee')
    textline(se,45,'Esta matriz explora inversión hipotética total. No es cotización ni propuesta de fee. No produce ROI sin Finance.')
    textline(se,46,'75% relevancia/eligibilidad es supuesto de prueba común. Ajustar Supuestos D13 cambia todas las visitas y cohortes.')
    header(se,49,['Aceptación geo pool amplio','Demanda adicional / mes','Fondeadas extra 6m'])
    for row,share in [(50,0),(51,.5),(52,1)]:
        value(se,f'C{row}',share,PCT,input_cell=True)
        value(se,f'D{row}',f'=IF(COUNT(\'Supuestos\'!$D$13)=1,1600*C{row}*\'Supuestos\'!$D$13,"N/D")',NUM)
        value(se,f'E{row}',f'=IF(COUNT(D{row},\'Supuestos\'!$F$23:$K$23,\'Supuestos\'!$F$33,\'Supuestos\'!$F$38)=9,D{row}*SUM(\'Supuestos\'!$F$23:$K$23)*\'Supuestos\'!$F$33*\'Supuestos\'!$F$38,"N/D")',NUM)
    textline(se,54,'Pool alto rendimiento 1.600 excluido de caso activo por SERP geográfica mixta. Esta prueba no lo vuelve elegible.')
    # Sources and assumptions ledger, kept out of outputs.
    f.column_dimensions['C'].width=18;f.column_dimensions['D'].width=32;f.column_dimensions['E'].width=60
    f.column_dimensions['F'].width=68;f.column_dimensions['G'].width=16;f.column_dimensions['H'].width=38
    f.print_area='C1:H35'
    header(f,7,['ID','Naturaleza','Fuente','Alcance','Fecha','Confianza'])
    srcs=cfg['sources']+[{'id':'DEMANDA','nature':cfg['demand_status'],'source':cfg['demand_source'],'scope':'Demanda mensual: cobertura y deduplicación deben corresponder al país/producto','date':cfg['as_of'],'confidence':'Provisional hasta validación de dataset'},
                        {'id':'RELEVANCIA','nature':'Supuesto ilustrativo','source':'Inputs editables, sin conversión observada','scope':cfg['relevance_status'],'date':cfg['as_of'],'confidence':'No inferencia empírica'},
                        {'id':'VALOR','nature':'Pendiente Finance','source':'Sin método bancario validado','scope':'Contribución neta anual por USD de saldo externo, método único','date':cfg['as_of'],'confidence':'No disponible'},
                        {'id':'COSTOS','nature':'Pendiente oferta','source':'Inputs editables, sin quote aprobado','scope':cfg.get('client_investment_status','Inversión pendiente')+' Costo interno Efeonce separado, pendiente.','date':cfg['as_of'],'confidence':'No inferencia de precio'}]
    for row,item in enumerate(srcs,8):
        for col,key in zip('CDEFGH',['id','nature','source','scope','date','confidence']):
            value(f,f'{col}{row}',item[key]);f[f'{col}{row}'].alignment=Alignment(wrap_text=True,vertical='top')
        f.row_dimensions[row].height=52
    textline(f,20,'Validación: Comercial confirma meta/horizonte; SEO deduplica demanda; banco aporta embudo; Finance valora fondeo.')
    textline(f,22,'Todos los drivers SIM son ejemplos aritméticos. No hay tasas medidas, costos aprobados ni causalidad probada.')
    dataset=json.loads((HERE/'demand-source.json').read_text())
    assert cfg['demand_monthly']==dataset['metadata']['default_envelope_searches_monthly'],'Source demand envelope changed: update inputs.json and narrative before regenerate'
    header(demand,7,['Grupo','Consulta representativa','Proxy mensual','Usar','Incluido'])
    demand.column_dimensions['C'].width=25;demand.column_dimensions['D'].width=56;demand.column_dimensions['E'].width=18;demand.column_dimensions['F'].width=12;demand.column_dimensions['G'].width=18
    for row,g in enumerate(dataset['groups'],8):
        for col,v in [('C',g['id']),('D',g['representative_keyword']),('E',g['monthly_volume_proxy']),('F','Sí' if g['use_in_default_envelope'] else 'No')]:value(demand,f'{col}{row}',v,NUM if col=='E' else None)
        value(demand,f'G{row}',f'=IF(F{row}="Sí",E{row},0)',NUM)
    value(demand,'C22','Proxy total mercado / mes');value(demand,'G22','=SUM(G8:G20)',NUM)
    textline(demand,24,'Agregación MAX por grupo editorial. Persiste solapamiento entre grupos y relevancia internacional; no personas únicas.')
    textline(demand,25,'PE/es. Consulta 06/10/2026. Volúmenes snapshot sept2026, últimos monthly_searches agosto. No baseline post-launch.')
    textline(demand,26,'Marca Pibank, competidores, soporte, educación, plazo fijo y futuros dólares excluidos. Ausencias preservadas NULL en dataset.')
    textline(demand,27,'Source: demand-source.json (metadata/grupos + SHA256 original). DataForSEO keyword_overview y keyword_suggestions live.')
    demand.print_area='C1:K28'
    # Summary depends on the owning builds, never on audit cells.
    textline(s,5,'ILUSTRACIÓN CONDICIONADA. Preparación interna. No es forecast aprobado ni compromiso de aperturas.')
    header(s,7,['Resultado de seis meses','Caso activo','Unidad'])
    output_rows={8:('Visitas adicionales',"=SUM('Mensual'!F13:K13)",'visitas',NUM),9:('Aperturas incrementales',"=SUM('Mensual'!F19:K19)",'cuentas',NUM),10:('Primer fondeo incremental',"=SUM('Mensual'!F23:K23)",'cuentas',NUM),11:('Entrada inicial bruta USD equivalente',"=D10*'Supuestos'!D10",'USD eq.',USD),12:('Saldo incremental al cierre',"='Mensual'!K25",'USD eq.',USD),13:('Saldo externo incremental al cierre',"='Mensual'!K26",'USD eq.',USD),14:('Escala / meta cuentas provisional',"=D10/'Supuestos'!D8",'%',PCT),16:('Costo completo programa',"='Supuestos'!D20",'USD',USD),17:('Costo / fondeada incremental','=IF(AND(\'Supuestos\'!D16="Sí",COUNT(D16)=1,D16>0,D10>0),D16/D10,"N/D")','USD',USD),18:('Valor mínimo / fondeada para equilibrio','=IF(AND(COUNT(D16)=1,D16>0,D10>0),D16/D10,"N/D")','USD en 6m',USD),19:('ROI seis meses',"='Mensual'!K31",'%',PCT),20:('Payback dentro de seis meses','=IF(COUNT(\'Mensual\'!F34:K34)=6,IFERROR(MATCH(1,\'Mensual\'!F34:K34,0),"No en 6m"),"N/D")','mes',NUM)}
    for row,(label,formula,unit,fmt) in output_rows.items():
        if row in (8,9,10):
            mr={8:13,9:19,10:23}[row];formula=f'=IF(COUNT(\'Mensual\'!F{mr}:K{mr})=6,SUM(\'Mensual\'!F{mr}:K{mr}),"N/D")'
        elif row==11:formula='=IF(COUNT(D10,\'Supuestos\'!D10)=2,D10*\'Supuestos\'!D10,"N/D")'
        elif row==14:formula='=IF(AND(COUNT(D10,\'Supuestos\'!D8)=2,\'Supuestos\'!D8>0),D10/\'Supuestos\'!D8,"N/D")'
        elif row==17:formula='=IF(AND(\'Supuestos\'!D16="Sí",COUNT(D16,D10)=2,D16>0,D10>0),D16/D10,"N/D")'
        elif row==18:formula='=IF(AND(COUNT(D16,D10)=2,D16>0,D10>0),D16/D10,"N/D")'
        value(s,f'C{row}',label);c=value(s,f'D{row}',formula,fmt);c.font=Font(name='Arial',size=10,color='18294A');value(s,f'E{row}',unit).alignment=Alignment(indent=1,vertical='center')
    s.column_dimensions['E'].width=15
    textline(s,22,'ROI y payback N/D hasta Finance, valor económico y costos completos validados. Saldo no es revenue.')
    header(s,24,['Mes','Visitas extra','Fondeadas extra','Saldo cierre USD eq.','Saldo externo USD eq.'])
    for row,col in enumerate(cols,25):
        for dest,src in [('C',8),('D',13),('E',23),('F',25),('G',26)]:value(s,f'{dest}{row}',f"='Mensual'!{col}{src}",'mmm-yy' if dest=='C' else USD if dest in 'FG' else NUM).font=Font(name='Arial',size=10,color='18294A')
    textline(s,32,'Inicio enero 2027 y seis meses son calendario de prueba. Meta USD inferida; moneda y periodo sin confirmación.')
    textline(s,33,'Producto público Soles. USD es equivalencia aritmética; no supone cuenta en dólares ni tipo de cambio.')
    textline(s,34,'Demanda 3.900 es proxy histórico, no forecast 2027. Relevancia 75%, referencia, captura y embudo son ilustrativos.')
    textline(s,36,'Cambiar selector en Supuestos D4 para recalcular Conservador/Base/Favorable con el mismo build.')
    for ws in wb:
        for row in ws:
            for c in row:
                if c.value is not None and c.font.name!='Arial':c.font=Font(name='Arial',size=10,color='18294A')
        ws.sheet_properties.outlinePr.summaryRight=False
    wb.save(output)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--no-copy',action='store_true');args=parser.parse_args()
    cfg=json.loads((HERE/'inputs.json').read_text());validate_inputs(cfg)
    original=HERE.parent/'research/pibank-demand-2026-10-06/demand-inputs.json'
    if original.exists():
        import hashlib
        raw=original.read_bytes();data=json.loads(raw)
        subset={'metadata':data['metadata'],'groups':data['groups'],'original_file':str(original),
                'original_sha256':hashlib.sha256(raw).hexdigest(),'scope':'Metadata/grupos de fuente, no listado completo de keywords'}
        (HERE/'demand-source.json').write_text(json.dumps(subset,ensure_ascii=False,indent=2)+'\n')
    with tempfile.TemporaryDirectory(prefix='pibank-author-') as tmp:
        raw=Path(tmp)/'raw';raw.mkdir();made=raw/NAME;make_book(cfg,made)
        final=recalculate(made,Path(tmp)/'calculated');shutil.copy2(final,HERE/f'.{NAME}.building');(HERE/f'.{NAME}.building').replace(HERE/NAME)
    result={'as_of':cfg['as_of'],'scope':'Demanda proxy de estimaciones mercado históricas. Escenarios/conversiones ilustrativos; seis meses desde inicio editable',
            'goal_status':cfg['goal_status'],'demand_status':cfg['demand_status'],'demand_monthly':cfg['demand_monthly'],
            'relevance_eligibility_factor':cfg['relevance_eligibility_factor'],'client_investment_status':cfg['client_investment_status'],
            'selected_case':cfg['selected_case'],'cases':[simulate(cfg,i) for i in (1,2,3)],
            'recalculation_engine':'Bundled LibreOffice','workbook':NAME}
    base=cfg['cases'][1];base_output=result['cases'][1]
    result['scale_tests_base_rates']=[{'share_of_provisional_account_goal':share,'accounts_required':cfg['goal_accounts']*share,
                                     'visits_required':cfg['goal_accounts']*share/(base['opening_rate']*base['funding_rate']),
                                     'first_deposit_usd_equivalent':cfg['goal_accounts']*share*cfg['ticket_usd'],
                                     'nature':'Prueba de escala, no cuota prometida; horizonte meta no confirmado'} for share in (.01,.05,.10)]
    result['investment_sensitivity_base']=[{'hypothetical_total_investment_usd':x,'minimum_value_per_incremental_funded_account_usd':x/base_output['funded_incremental'],
                                           'nature':'Inversión hipotética total; no oferta, fee, ROI ni margen Finance validado'} for x in cfg['investment_sensitivity_usd']]
    result['delay_sensitivity_base']=[{'delay_months':d,'visits_within_six_months':sum(x['visits_incremental'] for x in base_output['monthly'][:6-d]),
                                      'funded_within_six_months':sum(x['funded_incremental'] for x in base_output['monthly'][:6-d])} for d in (0,1,2)]
    (HERE/'results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    if not args.no_copy:
        DELIVERY.mkdir(parents=True,exist_ok=True)
        for p in HERE.iterdir():
            if p.is_file() and p.suffix in ('.xlsx','.json','.md','.py'):shutil.copy2(p,DELIVERY/p.name)
    print(json.dumps({'workbook':str(HERE/NAME),'case_outputs':[{k:v for k,v in cs.items() if k not in ('monthly','cohort_balances')} for cs in result['cases']]},ensure_ascii=False))

if __name__=='__main__':main()
