"""Independent native-engine checks of formulas, caches, scenarios and representative edits."""
from pathlib import Path
import hashlib, json, math, shutil, subprocess, tempfile
import openpyxl

HERE=Path(__file__).resolve().parent
NAME='Pibank-caso-negocio-2027.xlsx'
SOFFICE=Path('/Users/jreye/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
ERRORS={'#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#NULL!','#SPILL!','#CALC!'}

def checkclose(actual,expected,label):
    if not isinstance(actual,(int,float)) or not math.isclose(actual,expected,rel_tol=1e-9,abs_tol=1e-6):
        raise AssertionError(f'{label}: {actual!r} != {expected!r}')

def independent(wb):
    a=wb['Supuestos'];m=wb['Mensual'];s=wb['Resumen'];co=wb['Cohortes'];se=wb['Sensibilidad'];cols=list('FGHIJK')
    for ws in wb:
        for row in ws:
            for cell in row:
                assert not (cell.data_type=='e' or isinstance(cell.value,str) and cell.value in ERRORS),(ws.title,cell.coordinate,cell.value)
    critical=[a['D13'].value]+[a[f'{c}{r}'].value for c in cols for r in (23,28,33,38,43)]
    if not all(isinstance(x,(int,float)) for x in critical):
        assert s['D10'].value=='N/D' and s['D12'].value=='N/D'
        assert s['D19'].value=='N/D' and s['D20'].value=='N/D'
        return {'case':a['D5'].value,'funded_incremental':'N/D','balance_end_usd_equivalent':'N/D','roi':'N/D','payback':'N/D'}
    visits=[];opened=[];funded=[];stocks=[];external=[]
    for col in cols:
        demand=a[f'{col}43'].value;capture=a[f'{col}23'].value;op=a[f'{col}33'].value;fu=a[f'{col}38'].value
        v=demand*a['D13'].value*capture;o=v*op;n=o*fu
        visits.append(v);opened.append(o);funded.append(n)
        checkclose(m[f'{col}13'].value,v,f'visits {col}')
        checkclose(m[f'{col}23'].value,n,f'funded {col}')
        checkclose(m[f'{col}22'].value-m[f'{col}21'].value,n,f'two trajectories {col}')
    ticket=a['D10'].value
    for j,col in enumerate(cols):
        saldo=0
        for i in range(j+1):
            origin=cols[i];age=j-i
            relative=1 if age==0 else a[f'{origin}{48 if age==1 else 53 if age==2 else 58}'].value
            cohort=funded[i]*ticket*relative
            checkclose(co.cell(8+i,11+j).value,cohort,f'cohort {i}/{j}')
            saldo+=cohort
        stocks.append(saldo);external.append(saldo*a[f'{col}63'].value)
        checkclose(m[f'{col}25'].value,saldo,f'stock {col}')
        checkclose(m[f'{col}26'].value,external[-1],f'external {col}')
    for cell,expected in {'D8':sum(visits),'D9':sum(opened),'D10':sum(funded),'D11':sum(funded)*ticket,'D12':stocks[-1],'D13':external[-1],'D14':sum(funded)/a['D8'].value}.items():checkclose(s[cell].value,expected,cell)
    # Genuine self-contained sensitivity checks, including non-base combinations.
    for r in (17,19,21):
        for col in ('F','H','J'):
            expected=sum(visits)*se[f'C{r}'].value*se[f'{col}16'].value
            checkclose(se[f'{col}{r}'].value,expected,f'sensitivity {col}{r}')
    for row in (25,26,27):
        delay=se[f'C{row}'].value
        checkclose(se[f'D{row}'].value,sum(visits[:6-delay]),f'delay {delay}')
    checkclose(se['D33'].value,stocks[-1]*.4,'external sensitivity')
    investment=[a[f'{c}19'].value for c in cols]
    complete=all(isinstance(x,(int,float)) for x in investment)
    total=sum(investment) if complete else None
    finance=a['D14'].value=='Sí' and a['D16'].value=='Sí' and isinstance(a['D15'].value,(int,float)) and complete and total>0
    if finance:
        values=[x*a['D15'].value/12 for x in external]
        checkclose(s['D19'].value,(sum(values)-total)/total,'ROI')
        payback=next((i+1 for i in range(6) if sum(values[:i+1])>=sum(investment[:i+1])), 'No en 6m')
        assert s['D20'].value==payback,(s['D20'].value,payback)
    else:
        assert s['D19'].value=='N/D' and s['D20'].value=='N/D', (s['D19'].value,s['D20'].value)
    for ws in wb:
        for row in ws:
            for cell in row:
                assert not (cell.data_type=='e' or isinstance(cell.value,str) and cell.value in ERRORS),(ws.title,cell.coordinate,cell.value)
    assert co['E27'].value=='Fuera 6m' and co['G25'].value=='Fuera 6m'
    return {'case':a['D5'].value,'funded_incremental':sum(funded),'balance_end_usd_equivalent':stocks[-1],'roi':s['D19'].value,'payback':s['D20'].value}

def recalc(path,out):
    out.mkdir()
    with tempfile.TemporaryDirectory(prefix='pibank-validation-profile-') as profile:
        p=subprocess.run([str(SOFFICE),f'-env:UserInstallation={Path(profile).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=90)
        assert p.returncode==0 and (out/path.name).exists(),(p.returncode,p.stdout,p.stderr)
    return openpyxl.load_workbook(out/path.name,data_only=True)

def main():
    source=HERE/NAME;cases=[]
    saved=openpyxl.load_workbook(source,data_only=True)
    cases.append({'test':'Delivered cached values','result':independent(saved)})
    formulas=openpyxl.load_workbook(source,data_only=False)
    assert formulas['Resumen']['D10'].data_type=='f'
    assert formulas['Mensual']['K23'].data_type=='f'
    assert formulas['Sensibilidad']['H19'].data_type=='f'
    assert formulas['Supuestos']['D4'].value==2
    assert formulas['Supuestos']['D21'].value is None
    modifications=[('Conservador',{'D4':1}),('Favorable',{'D4':3}),('Edit later-period funding',{'K40':.325}),
                   ('Finance validation alone',{'D14':'Sí','D15':.08}),
                   ('Finance and investment validated',{'D14':'Sí','D16':'Sí','D15':.08,'F19':1000,'G19':1000,'H19':1000,'I19':1000,'J19':1000,'K19':1000}),
                   ('Missing investment',{'K19':None,'D14':'Sí','D16':'Sí','D15':.08}),
                   ('Zero investment',{'F19':0,'G19':0,'H19':0,'I19':0,'J19':0,'K19':0,'D14':'Sí','D16':'Sí','D15':.08}),
                   ('Blank unselected case',{'F24':None}),('Blank selected capture',{'F25':None}),
                   ('Blank relevance',{'D13':None}),('Zero relevance',{'D13':0})]
    with tempfile.TemporaryDirectory(prefix='pibank-validation-') as tmp:
        for i,(name,changes) in enumerate(modifications):
            wb=openpyxl.load_workbook(source);ws=wb['Supuestos']
            for cell,v in changes.items():ws[cell]=v
            path=Path(tmp)/f'test-{i}.xlsx';wb.save(path)
            recalced=recalc(path,Path(tmp)/f'calc-{i}')
            result=independent(recalced)
            cases.append({'test':name,'changed_inputs':changes,'result':result})
    assert cases[1]['result']['funded_incremental']<cases[0]['result']['funded_incremental']<cases[2]['result']['funded_incremental']
    assert cases[3]['result']['funded_incremental']<cases[0]['result']['funded_incremental']
    cfg=json.loads((HERE/'inputs.json').read_text());outputs=json.loads((HERE/'results.json').read_text())
    # Independently reconcile generated JSON against actual recalculated selected/case outputs.
    for checked in cases[:3]:
        matched=next(cs for cs in outputs['cases'] if cs['case']==checked['result']['case'])
        checkclose(matched['funded_incremental'],checked['result']['funded_incremental'],'JSON funded')
        checkclose(matched['balance_end_usd_equivalent'],checked['result']['balance_end_usd_equivalent'],'JSON stock')
    result={'status':'passed','engine':'Bundled LibreOffice','workbook_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
            'checks':cases,'unverified':'Excel UI interaction not performed; formulas/control recalculated in LibreOffice','source_inputs_finance_validated':cfg['finance_validated']}
    (HERE/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    delivery=Path('/Users/jreye/Documents/Banco Pichincha Peru — Prospect Case/02-Uso-interno/modelo')
    if delivery.exists():shutil.copy2(HERE/'validation.json',delivery/'validation.json')
    print(json.dumps({'status':'passed','native_recalculation_tests':len(cases),'finance_outputs_delivered':cases[0]['result']},ensure_ascii=False))

if __name__=='__main__':main()
