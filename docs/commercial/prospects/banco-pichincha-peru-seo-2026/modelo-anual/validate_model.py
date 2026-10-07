"""Independent checks against native recalculated XLSX, without importing the generator."""
from pathlib import Path
from datetime import date, datetime
import calendar, hashlib, json, math, shutil, subprocess, tempfile
import openpyxl

HERE=Path(__file__).resolve().parent
NAME='Pibank-caso-negocio-anual-2027.xlsx'
SOFFICE=Path('/Users/jreye/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
COLS=[openpyxl.utils.get_column_letter(i) for i in range(6,18)]
CCOLS=[openpyxl.utils.get_column_letter(i) for i in range(12,24)]

def close(x,y,label):
    assert isinstance(x,(int,float)) and math.isclose(x,y,rel_tol=1e-9,abs_tol=1e-6),(label,x,y)

def relative(days,ys):
    if days<0:return 0.0
    xs=[0,30,60,90,180,365];ys=[1]+ys
    for i in range(1,len(xs)):
        if days<=xs[i]:return ys[i-1]+(ys[i]-ys[i-1])*(days-xs[i-1])/(xs[i]-xs[i-1])
    return ys[-1]

def shift(d,n,day):
    y,m=divmod(d.year*12+d.month-1+n,12)
    return date(y,m+1,day)

def independent(wb):
    for ws in wb:
        for row in ws:
            for c in row:assert c.data_type!='e',(ws.title,c.coordinate,c.value)
    a=wb['Supuestos'];m=wb['Mensual'];co=wb['Cohortes'];s=wb['Resumen'];se=wb['Sensibilidad']
    critical=[a['D13'].value,a['D18'].value,a['D19'].value,a['D25'].value]+[a[f'{c}{r}'].value for c in COLS for r in (28,30,40,45,50)]
    if not all(isinstance(x,(int,float)) for x in critical):
        assert s['D10'].value=='N/D' and s['D13'].value=='N/D'
        assert s['D20'].value=='N/D' and s['D21'].value=='N/D'
        return {'case':a['D5'].value,'funded_incremental':'N/D','balance_end2027_usd_equivalent':'N/D','roi':'N/D','payback':'N/D'}
    dates=[a[f'{c}21'].value.date() for c in COLS]
    ends=[d.replace(day=calendar.monthrange(d.year,d.month)[1]) for d in dates]
    visits=[];opens=[];potential=[];funddates=[]
    for c,d in zip(COLS,dates):
        v=a[f'{c}50'].value*a[f'{c}28'].value*a['D25'].value*a['D13'].value*a[f'{c}30'].value
        o=v*a[f'{c}40'].value;n=o*a[f'{c}45'].value
        visits.append(v);opens.append(o);potential.append(n);funddates.append(shift(d,a['D18'].value,a['D19'].value))
        close(m[f'{c}16'].value,v,'visits'+c);close(m[f'{c}21'].value,o,'openings'+c)
        close(m[f'{c}25'].value,n,'potential'+c)
        close(m[f'{c}24'].value-m[f'{c}23'].value,n,'two trajectories'+c)
    funded=[sum(n for n,fd in zip(potential,funddates) if fd.year==d.year and fd.month==d.month) for d in dates]
    stocks=[0.]*12;extstocks=[0.]*12;average=[0.]*12;extaverage=[0.]*12
    for i,(c,n,fd) in enumerate(zip(COLS,potential,funddates)):
        origin=a[f'{c}80'].value;knots=[a[f'{c}{r}'].value for r in (55,60,65,70,75)];initial=n*a['D10'].value;prior=0
        assert co[f'C{9+i}'].value.date()==fd
        for j,(cc,start,end) in enumerate(zip(CCOLS,dates,ends)):
            amount=initial*relative((end-fd).days,knots)
            mean=0 if fd>end else ((initial+amount)/2*((end-fd).days+1)/(end-start).days.__add__(1) if fd>=start else (prior+amount)/2)
            close(co[f'{cc}{9+i}'].value,amount,f'cohort {i}/{j}')
            close(co[f'{cc}{30+i}'].value,amount*origin,f'external origin {i}/{j}')
            close(co[f'{cc}{50+i}'].value,mean,f'average {i}/{j}')
            close(co[f'{cc}{70+i}'].value,mean*origin,f'average external {i}/{j}')
            stocks[j]+=amount;extstocks[j]+=amount*origin;average[j]+=mean;extaverage[j]+=mean*origin;prior=amount
        for dest,days,k in (('E',30,0),('F',60,1),('G',90,2),('H',180,3),('I',365,4)):
            v=co[f'{dest}{88+i}'].value
            if (ends[-1]-fd).days<days:assert v=='Fuera2027',(dest,i,v)
            else:close(v,initial*knots[k],f'maturity {days}/{i}')
    for j,c in enumerate(COLS):
        for row,value in ((28,funded[j]),(35,stocks[j]),(36,extstocks[j]),(37,average[j]),(38,extaverage[j])):close(m[f'{c}{row}'].value,value,f'month {c}{row}')
    checks={'D8':sum(visits),'D9':sum(opens),'D10':sum(funded),'D11':sum(potential)-sum(funded),'D12':sum(funded)*a['D10'].value,'D13':stocks[-1],'D14':extstocks[-1],'D15':sum(extaverage)/12,'D16':sum(funded)/a['D8'].value}
    for cell,value in checks.items():close(s[cell].value,value,cell)
    for row in (16,18,20):
        for c in ('F','H','J'):close(se[f'{c}{row}'].value,sum(visits)*se[f'C{row}'].value*se[f'{c}15'].value,'sensitivity')
    for row in (24,25,26):
        lag=se[f'C{row}'].value;close(se[f'D{row}'].value,sum(potential[:12-lag]),'lag sensitivity')
    costs=[a[f'{c}22'].value for c in COLS]
    ready=a['D14'].value=='Sí' and a['D16'].value=='Sí' and isinstance(a['D15'].value,(int,float)) and all(isinstance(v,(int,float)) and v>=0 for v in costs) and sum(costs)>0
    if ready:
        values=[v*a['D15'].value/12 for v in extaverage]
        close(s['D20'].value,(sum(values)-sum(costs))/sum(costs),'ROI within horizon')
        for c,v in zip(COLS,values):close(m[f'{c}42'].value,v,'monthly Finance')
        payback=next((j+1 for j in range(12) if sum(values[:j+1])>=sum(costs[:j+1])), 'No en2027')
        assert s['D21'].value==payback,(s['D21'].value,payback)
    else:assert s['D20'].value=='N/D' and s['D21'].value=='N/D',(s['D20'].value,s['D21'].value)
    return {'case':a['D5'].value,'funded_incremental':sum(funded),'funded_scheduled_after2027':sum(potential)-sum(funded),'balance_end2027_usd_equivalent':stocks[-1],'external_balance_end2027_usd_equivalent':extstocks[-1],'average_external_balance_2027_usd_equivalent':sum(extaverage)/12,'roi':s['D20'].value,'payback':s['D21'].value}

def native(path,out):
    out.mkdir()
    with tempfile.TemporaryDirectory(prefix='annual-profile-') as profile:
        p=subprocess.run([str(SOFFICE),f'-env:UserInstallation={Path(profile).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=90)
        assert p.returncode==0 and (out/path.name).exists(),(p.stdout,p.stderr)
    return openpyxl.load_workbook(out/path.name,data_only=True)

def main():
    path=HERE/NAME;saved=openpyxl.load_workbook(path,data_only=True);checks=[{'test':'Delivered cached Base,12 months, cohorts, maturity and sensitivity','result':independent(saved)}]
    formulas=openpyxl.load_workbook(path,data_only=False)
    assert formulas['Resumen']['D10'].data_type=='f' and formulas['Cohortes']['W20'].data_type=='f'
    assert formulas['Supuestos']['D4'].value==2 and formulas['Supuestos']['D15'].value is None and formulas['Supuestos']['Q22'].value is None
    history=[saved['Historia'][f'{c}9'].value for c in COLS]
    financed={'D14':'Sí','D16':'Sí','D15':.08,**{f'{c}22':1000 for c in COLS}}
    changes=[('Conservador',{'D4':1}),('Favorable',{'D4':3}),('December funding lower',{'Q47':.325}),
        ('Finance only keeps ROI N/D',{'D14':'Sí','D15':.08}),('Finance and complete investment',financed),
        ('Missing December expense',{**financed,'Q22':None}),('Zero expense not unknown',{**financed,**{f'{c}22':0 for c in COLS}}),
        ('Blank unselected driver',{'F31':None}),('Blank selected driver',{'F32':None}),('Blank eligibility',{'D13':None}),('Zero eligibility',{'D13':0}),
        ('Funding lag1',{**financed,'D18':1}),('Funding lag2',{'D18':2}),('Funding lag12 excludes all2027',{'D18':12}),
        ('December origin external30pct sticks to cohort',{'Q82':.3}),('Funding day1',{**financed,'D19':1}),('Funding day28',{**financed,'D19':28}),
        ('Observed historical shape as hypothesis',{'D26':'Historia'}),('Zero December expense is measured zero',{**financed,'Q22':0}),('Negative expense rejected by financial gates',{**financed,'Q22':-1000}),('Validated zero net value',{**financed,'D15':0})]
    with tempfile.TemporaryDirectory(prefix='annual-validation-') as tmp:
        for i,(name,edits) in enumerate(changes):
            w=openpyxl.load_workbook(path)
            for c,v in edits.items():w['Supuestos'][c]=v
            p=Path(tmp)/f'test{i}.xlsx';w.save(p);r=native(p,Path(tmp)/f'calc{i}')
            assert [r['Historia'][f'{c}9'].value for c in COLS]==history,'Source history must remain immutable under forecast controls'
            checks.append({'test':name,'changed_inputs':edits,'result':independent(r)})
    assert checks[1]['result']['funded_incremental']<checks[0]['result']['funded_incremental']<checks[2]['result']['funded_incremental']
    assert checks[3]['result']['funded_incremental']<checks[0]['result']['funded_incremental']
    assert checks[14]['result']['funded_incremental']==0 and checks[14]['result']['balance_end2027_usd_equivalent']==0
    assert checks[16]['result']['average_external_balance_2027_usd_equivalent']>checks[17]['result']['average_external_balance_2027_usd_equivalent']
    outputs=json.loads((HERE/'results.json').read_text())
    for result in [c['result'] for c in checks[:3]]:
        target=next(c for c in outputs['cases'] if c['case']==result['case'])
        for key in ('funded_incremental','balance_end2027_usd_equivalent','external_balance_end2027_usd_equivalent','average_external_balance_2027_usd_equivalent'):close(target[key],result[key],'JSON '+key)
    record={'status':'passed','engine':'Bundled LibreOffice','workbook_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'native_recalculation_tests':len(checks),'checks':checks,'unverified':'Excel desktop UI not operated; native control/formula/cache changes tested in LibreOffice.'}
    (HERE/'validation.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'status':'passed','tests':len(checks),'selected_outputs':checks[0]['result']},ensure_ascii=False))

if __name__=='__main__':main()
