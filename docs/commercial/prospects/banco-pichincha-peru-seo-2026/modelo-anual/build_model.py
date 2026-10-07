"""Pibank annual editable model. Inputs and projections remain explicitly conditional.

Author with bundled openpyxl as requested. LibreOffice recalculates formulas and caches.
One selected-case build, source history separate from 2027 assumptions, preserved v6 original.
"""
from pathlib import Path
from datetime import date, timedelta
import argparse, calendar, hashlib, json, re, shutil, subprocess, tempfile
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

HERE=Path(__file__).resolve().parent
DELIVERY=Path('/Users/jreye/Documents/Banco Pichincha Peru — Prospect Case/02-Uso-interno/propuesta-2027-anual/modelo')
SOFFICE=Path('/Users/jreye/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
NAME='Pibank-caso-negocio-anual-2027.xlsx'
COLS=[openpyxl.utils.get_column_letter(n) for n in range(6,18)]
CCOLS=[openpyxl.utils.get_column_letter(n) for n in range(12,24)]
NUM='#,##0.0;(#,##0.0);"—"';USD='"US$"#,##0;("US$"#,##0);"—"';PCT='0.0%;(0.0%);"—"';RATE='0.00%;(0.00%);"—"'
DRIVERS={'extra_capture':30,'reference_capture':35,'opening_rate':40,'funding_rate':45,'demand':50,'retention30':55,'retention60':60,'retention90':65,'retention180':70,'retention365':75,'external_share':80}

def monthdate(start,offset):
    n=start.year*12+start.month-1+offset
    return date(n//12,n%12+1,1)

def monthend(d):return date(d.year,d.month,calendar.monthrange(d.year,d.month)[1])

def retention(age,knots):
    if age<0:return 0.0
    xs=[0,30,60,90,180,365];ys=[1.0]+knots
    for i in range(1,len(xs)):
        if age<=xs[i]:return ys[i-1]+(ys[i]-ys[i-1])*(age-xs[i-1])/(xs[i]-xs[i-1])
    return ys[-1]

def simulate(cfg,index):
    cs=cfg['cases'][index-1];start=date.fromisoformat(cfg['start_month']);periods=[monthdate(start,i) for i in range(12)];lag=cfg['funding_lag_months']
    seasonal=[1.0]*12 if cfg['seasonality_mode']=='Plano' else cfg['historical_seasonality_by_calendar_month']
    demand=[cfg['demand_monthly_proxy']*seasonal[p.month-1]*cfg['growth_multiplier']*cfg['relevance_eligibility_factor'] for p in periods]
    extra=[d*v for d,v in zip(demand,cs['extra_capture'])];opening=[v*cs['opening_rate'] for v in extra];potential=[o*cs['funding_rate'] for o in opening]
    funddates=[monthdate(p,lag).replace(day=cfg['funding_day_of_month']) for p in periods]
    cash=[sum(n for n,fd in zip(potential,funddates) if fd.year==p.year and fd.month==p.month) for p in periods]
    closes=[];avg=[];cohort_closes=[];cohort_avg=[];knots=[cs[k] for k in ('retention30','retention60','retention90','retention180','retention365')]
    for n,fd in zip(potential,funddates):
        close=[];av=[]
        for i,p in enumerate(periods):
            end=monthend(p);bal=n*cfg['ticket_usd_equivalent']*retention((end-fd).days,knots)
            close.append(bal)
            if fd>end:av.append(0.0)
            elif fd>=p:
                weight=(end-fd).days+1
                av.append((n*cfg['ticket_usd_equivalent']+bal)/2*weight/calendar.monthrange(p.year,p.month)[1])
            else:av.append(((close[i-1] if i else 0)+bal)/2)
        cohort_closes.append(close);cohort_avg.append(av)
    closes=[sum(c[j] for c in cohort_closes) for j in range(12)]
    avg=[sum(c[j] for c in cohort_avg) for j in range(12)]
    costs=cfg['monthly_client_investment_usd'];margin=cfg['net_annual_value_per_usd_external_average_balance']
    ready=cfg['finance_validated'] and cfg['investment_validated'] and isinstance(margin,(int,float)) and all(isinstance(x,(int,float)) and x>=0 for x in costs) and sum(costs)>0
    values=[x*cs['external_share']*margin/12 for x in avg] if ready else None
    payback=next((i+1 for i in range(12) if sum(values[:i+1])>=sum(costs[:i+1])),None) if ready else None
    return {'case':cs['name'],'nature':'Hipótesis ilustrativa anual, no forecast bancario validado','visits_incremental':sum(extra),'openings_incremental':sum(opening),
            'funded_incremental_within2027':sum(cash),'funded_scheduled_after2027':sum(potential)-sum(cash),'first_deposit_usd_equivalent_within2027':sum(cash)*cfg['ticket_usd_equivalent'],
            'balance_end2027_usd_equivalent':closes[-1],'external_balance_end2027_usd_equivalent':closes[-1]*cs['external_share'],
            'average_external_balance_2027_usd_equivalent':sum(avg)/12*cs['external_share'],'account_goal_share_arithmetic':sum(cash)/cfg['goal_accounts'],
            'roi':(sum(values)-sum(costs))/sum(costs) if ready else None,'payback_month':payback,'financial_status':'N/D hasta inversión y Finance validados' if not ready else 'Calculado sobre inputs validados',
            'monthly':[{'month':p.isoformat(),'visits_incremental':extra[i],'openings_incremental':opening[i],'funded_incremental':cash[i],
                        'balance_close_usd_equivalent':closes[i],'external_balance_close_usd_equivalent':closes[i]*cs['external_share'],
                        'average_external_balance_usd_equivalent':avg[i]*cs['external_share']} for i,p in enumerate(periods)]}

def put(ws,address,v,fmt=None,editable=False):
    if isinstance(v,str) and not v.startswith('=') and not v.startswith('http') and not v.startswith('../'):
        v=re.sub(r'(?<=[A-Za-zÁÉÍÓÚáéíóú])(?=\d)|(?<=\d)(?=[A-Za-zÁÉÍÓÚáéíóú])',' ',v)
        v=v.replace('N/Dhasta','N/D hasta').replace('Primerfondeo','Primer fondeo').replace('primerfondeo','primer fondeo').replace('dentro 2027','dentro de 2027').replace('después 2027','después de 2027').replace('origenexterno','origen externo').replace('costeEfeonce','coste Efeonce').replace('costoEfeonce','costo Efeonce').replace('Inversióncliente','Inversión cliente').replace('inversióncliente','inversión cliente').replace('saldomantenido','saldo mantenido')
    c=ws[address];c.value=v
    color='0000CC' if editable else '008050' if isinstance(v,str) and v.startswith('=') and '!' in v else '18294A'
    c.font=Font(name='Arial',size=10,color=color)
    c.alignment=Alignment(vertical='center')
    if fmt:
        c.number_format=fmt
        c.alignment=Alignment(horizontal='right',indent=1,vertical='center')
    if editable:c.fill=PatternFill('solid',fgColor='FFF2CC')
    return c

def note(ws,row,text,endcol=17):
    ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=endcol)
    put(ws,f'C{row}',text).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row].height=30

def head(ws,row,labels,start=3):
    for col,label in enumerate(labels,start):
        c=put(ws,f'{openpyxl.utils.get_column_letter(col)}{row}',label)
        c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='20385F')
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[row].height=30

def sheet(wb,name,title,lastrow=45,endcol=17):
    ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=ws.column_dimensions['B'].width=2
    ws.column_dimensions['C'].width=39;ws.column_dimensions['D'].width=17;ws.column_dimensions['E'].width=3
    for n in range(6,endcol+1):ws.column_dimensions[openpyxl.utils.get_column_letter(n)].width=14
    put(ws,'C2',title).font=Font(name='Arial',size=14,bold=True,color='18294A');ws.row_dimensions[2].height=25
    put(ws,'C4','Escenario seleccionado');put(ws,'D4',"='Supuestos'!D5")
    ws['D4'].alignment=Alignment(horizontal='center',vertical='center');ws['D4'].border=Border(bottom=Side(style='dotted',color='7E8A9E'))
    for r in range(5,lastrow+1):ws.row_dimensions[r].height=20
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_area=f'C1:{openpyxl.utils.get_column_letter(endcol)}{lastrow}';ws.print_options.horizontalCentered=True
    ws.page_margins.left=ws.page_margins.right=.2;ws.page_margins.top=ws.page_margins.bottom=.25
    return ws

def product_guard(parts):return f'IF(COUNT({",".join(parts)})={len(parts)},{"*".join(parts)},"N/D")'

def build(cfg,source,output):
    wb=openpyxl.Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    s=sheet(wb,'Resumen','Pibank: programa anual 2027',49,11)
    a=sheet(wb,'Supuestos','Controles y drivers 2027',92)
    m=sheet(wb,'Mensual','Adquisición, fondeo y valor mensual',51)
    co=sheet(wb,'Cohortes','Cohortes: saldo, origen externo y maduración',145,23)
    se=sheet(wb,'Sensibilidad','Escala y sensibilidad anual',46,11)
    raw=sheet(wb,'Demanda','Demanda de mercado: fuente y cobertura',max(42,len(source['groups'])+len(source['metadata']['limitations'])+22),11)
    hist=sheet(wb,'Historia','Serie histórica separada de hipótesis 2027',29)
    f=sheet(wb,'Fuentes','Ledger de evidencia y supuestos',40,11)
    for ws in (s,m,se,raw,hist,f):ws.page_setup.fitToHeight=1
    a.page_setup.fitToHeight=2
    co.page_setup.fitToHeight=3
    s.sheet_properties.tabColor='20385F';a.sheet_properties.tabColor='687DA0'
    for ws in (a,m,co):ws.freeze_panes='F9'
    # One editable scenario control, no parallel case builds.
    put(a,'C4','Selector escenario (1 / 2 / 3)');put(a,'D4',cfg['selected_case'],editable=True)
    put(a,'C5','Caso activo');put(a,'D5','=CHOOSE(D4,"Conservador","Base","Favorable")')
    note(a,6,'1 Conservador, 2 Base, 3 Favorable. Casos de prueba; nombres no representan probabilidades.')
    selector=DataValidation(type='whole',operator='between',formula1=1,formula2=3);selector.showErrorMessage=True;a.add_data_validation(selector);selector.add(a['D4'])
    global_rows=[(8,'Meta cuentas provisional',cfg['goal_accounts'],NUM),(9,'Meta USD provisional',cfg['goal_usd'],USD),(10,'Ticket USD equivalente provisional',cfg['ticket_usd_equivalent'],USD),
                 (11,'Inicio ilustrativo',date.fromisoformat(cfg['start_month']),'mmm-yy'),(13,'Relevancia/elegibilidad ilustrativa',cfg['relevance_eligibility_factor'],PCT),
                 (14,'Finance validó método neto','Sí' if cfg['finance_validated'] else 'No',None),(15,'Valor neto anual / USD saldo externo',cfg['net_annual_value_per_usd_external_average_balance'],RATE),
                 (16,'Inversión completa validada','Sí' if cfg['investment_validated'] else 'No',None),(17,'Capacidad total con reserva (horas)',cfg['delivery_capacity_hours_total'],NUM),
                 (18,'Lag apertura a fondeo (meses)',cfg['funding_lag_months'],'0'),(19,'Día fondeo supuesto (1–28)',cfg['funding_day_of_month'],'0'),
                 (25,'Multiplicador crecimiento2027 supuesto',cfg['growth_multiplier'],'0.00'),(26,'Estacionalidad usada',cfg['seasonality_mode'],None)]
    for r,label,v,fmt in global_rows:put(a,f'C{r}',label);put(a,f'D{r}',v,fmt,editable=r!=17)
    note(a,12,'US$38m es inferencia provisional. Programa anual2027 no confirma periodo ni definición de la meta bancaria.')
    for r in (14,16):
        dv=DataValidation(type='list',formula1='"Sí,No"');a.add_data_validation(dv);dv.add(a[f'D{r}'])
    for r,lo,hi in ((18,0,12),(19,1,28)):
        dv=DataValidation(type='whole',operator='between',formula1=lo,formula2=hi);dv.showErrorMessage=True;a.add_data_validation(dv);dv.add(a[f'D{r}'])
    mode=DataValidation(type='list',formula1='"Plano,Historia"');a.add_data_validation(mode);mode.add(a['D26'])
    put(a,'C20','Proxy mensual de fuente histórica');put(a,'D20',"='Demanda'!J5",NUM)
    for i,col in enumerate(COLS):put(a,f'{col}21',f'=EDATE($D$11,{i})','mmm-yy')
    put(a,'C22','Inversión cliente por mes USD (pendiente)')
    for i,col in enumerate(COLS):put(a,f'{col}22',cfg['monthly_client_investment_usd'][i],USD,editable=True)
    nonnegative=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0,allow_blank=True);nonnegative.showErrorMessage=True;nonnegative.errorStyle='stop';a.add_data_validation(nonnegative);nonnegative.add('F22:Q22')
    put(a,'C23','Inversión anual si completa');put(a,'D23','=IF(COUNT(F22:Q22)=12,IF(MIN(F22:Q22)>=0,SUM(F22:Q22),"N/D"),"N/D")',USD)
    put(a,'C24','Coste interno Efeonce USD (pendiente)');put(a,'D24',cfg['efeonce_internal_cost_usd'],USD,editable=True)
    put(a,'C27','Factor histórico observado / promedio');put(a,'C28','Factor usado en hipótesis2027')
    for i,col in enumerate(COLS):
        put(a,f'{col}27',f'=IF(COUNT(\'Historia\'!{col}11)=1,\'Historia\'!{col}11,"N/D")','0.00')
        put(a,f'{col}28',f'=IF($D$26="Plano",1,{col}27)','0.00')
    names={'extra_capture':'Captura adicional','reference_capture':'Captura referencia','opening_rate':'Visita a apertura','funding_rate':'Apertura a fondeo','demand':'Proxy demanda trasladado2027',
           'retention30':'Saldo relativo30d','retention60':'Saldo relativo60d','retention90':'Saldo relativo90d','retention180':'Saldo relativo180d','retention365':'Saldo relativo365d','external_share':'Origen externo de nueva cohorte'}
    for key,row in DRIVERS.items():
        put(a,f'C{row}',names[key]+' activo')
        for col in COLS:put(a,f'{col}{row}',f'=IF(CHOOSE($D$4,COUNT({col}{row+1}),COUNT({col}{row+2}),COUNT({col}{row+3}))=1,CHOOSE($D$4,{col}{row+1},{col}{row+2},{col}{row+3}),"N/D")',NUM if key=='demand' else RATE)
        for idx,case in enumerate(cfg['cases'],1):
            put(a,f'C{row+idx}',case['name']);vals=cfg['demand_monthly_proxy'] if key=='demand' else case[key]
            for i,col in enumerate(COLS):put(a,f'{col}{row+idx}',vals[i] if isinstance(vals,list) else vals,NUM if key=='demand' else RATE,editable=True)
        constraint=DataValidation(type='decimal',operator='greaterThanOrEqual' if key=='demand' else 'between',formula1=0,formula2=None if key=='demand' else 1,allow_blank=True);constraint.showErrorMessage=True;constraint.errorStyle='stop';a.add_data_validation(constraint);constraint.add(f'F{row+1}:Q{row+3}')
    note(a,85,'M1 establece accesos, medición y delivery. No evalúa viabilidad SEO por resultados tempranos ni activa un piloto.')
    note(a,86,'Programa12meses; revisiones3/6/9/12 para priorizar y optimizar. No prometen ROI temprano ni justifican corte a3meses.')
    note(a,87,'Origen externo se fija en cada cohorte desde el mes de apertura. No se repondera todo el saldo por la tasa del mes actual.')
    note(a,88,'Retenciones son hipótesis por cuenta original y contienen retiros/abandono. Se interpola entre puntos0/30/60/90/180/365d.')
    note(a,89,'Estacionalidad Historia reusa patrón pasado como hipótesis, no forecast2027. Plano=1 no afirma ausencia de estacionalidad.')
    note(a,90,'USD es equivalencia de meta/ticket; producto público Soles. Sin FX supuesto ni disponibilidad de cuenta USD inventada.')
    note(a,91,'Valor Finance es neto por UN método. No sumar FTP y ahorro de fondeo duplicados. Inversióncliente y costeEfeonce separados.')
    # Monthly build.
    head(m,7,['Indicador','Unidad','']+[f'Mes{i+1}' for i in range(12)])
    labels={9:('Proxy mensual histórico trasladado','búsquedas'),10:('Estacionalidad usada','factor'),11:('Crecimiento2027 supuesto','factor'),12:('Demanda tras relevancia','búsquedas'),13:('Captura referencia','%'),14:('Captura adicional programa','%'),15:('Visitas referencia','visitas'),16:('Visitas incrementales','visitas'),17:('Visitas programa','visitas'),18:('Visita a apertura','%'),19:('Aperturas referencia','cuentas'),20:('Aperturas programa','cuentas'),21:('Aperturas incrementales','cuentas'),22:('Apertura a fondeo','%'),23:('Fondeo potencial referencia','cuentas'),24:('Fondeo potencial programa','cuentas'),25:('Fondeo potencial incremental','cuentas'),26:('Primerfondeo referencia dentro año','cuentas'),27:('Primerfondeo programa dentro año','cuentas'),28:('Primerfondeo incremental dentro año','cuentas'),35:('Saldo incremental al cierre','USD eq.'),36:('Saldo externo al cierre','USD eq.'),37:('Saldo medio mensual incremental','USD eq.'),38:('Saldo medio mensual externo','USD eq.'),40:('Inversión cliente pendiente','USD'),41:('Inversión acumulada','USD'),42:('Valor neto mensual Finance','USD'),43:('Valor acumulado Finance','USD'),44:('ROI acumulado con gates Finance','%'),45:('Valor menos inversión acumulada','USD'),46:('Recuperación alcanzada','0/1')}
    for r,(label,unit) in labels.items():put(m,f'C{r}',label);put(m,f'D{r}',unit)
    for i,col in enumerate(COLS):
        put(m,f'{col}8',f"='Supuestos'!{col}21",'mmm-yy')
        for r,src in ((9,50),(10,28),(13,35),(14,30),(18,40),(22,45)):put(m,f'{col}{r}',f"='Supuestos'!{col}{src}",NUM if r==9 else RATE if r in (13,14,18,22) else '0.00')
        put(m,f'{col}11',"=IF(COUNT('Supuestos'!$D$25)=1,'Supuestos'!$D$25,\"N/D\")",'0.00')
        formulas={12:product_guard([f'{col}9',f'{col}10',f'{col}11',"'Supuestos'!$D$13"]),15:product_guard([f'{col}12',f'{col}13']),16:product_guard([f'{col}12',f'{col}14']),17:f'IF(COUNT({col}15:{col}16)=2,SUM({col}15:{col}16),"N/D")',19:product_guard([f'{col}15',f'{col}18']),20:product_guard([f'{col}17',f'{col}18']),21:f'IF(COUNT({col}19:{col}20)=2,{col}20-{col}19,"N/D")',23:product_guard([f'{col}19',f'{col}22']),24:product_guard([f'{col}20',f'{col}22']),25:f'IF(COUNT({col}23:{col}24)=2,{col}24-{col}23,"N/D")'}
        for r,form in formulas.items():put(m,f'{col}{r}','='+form,NUM)
        for target,src in ((26,23),(27,24)):
            form=f'=IF(COUNT(\'Supuestos\'!$D$18)=1,IF(\'Supuestos\'!$D$18<={i},INDEX($F${src}:$Q${src},1,{i+1}-\'Supuestos\'!$D$18),0),"N/D")'
            put(m,f'{col}{target}',form,NUM)
        put(m,f'{col}28',f'=IF(COUNT({col}26:{col}27)=2,{col}27-{col}26,"N/D")',NUM)
        for r,src in ((35,22),(36,43),(37,63),(38,83)):put(m,f'{col}{r}',f"='Cohortes'!{CCOLS[i]}{src}",USD)
        put(m,f'{col}40',f'=IF(COUNT(\'Supuestos\'!{col}22)=1,IF(\'Supuestos\'!{col}22>=0,\'Supuestos\'!{col}22,"N/D"),"N/D")',USD)
        put(m,f'{col}41',f'=IF(COUNT($F40:{col}40)={i+1},SUM($F40:{col}40),"N/D")',USD)
        put(m,f'{col}42',f'=IF(AND(\'Supuestos\'!$D$14="Sí",COUNT(\'Supuestos\'!$D$15,{col}38)=2),{col}38*\'Supuestos\'!$D$15/12,"N/D")',USD)
        put(m,f'{col}43',f'=IF(COUNT($F42:{col}42)={i+1},SUM($F42:{col}42),"N/D")',USD)
        gate=f'AND(\'Supuestos\'!$D$14="Sí",\'Supuestos\'!$D$16="Sí",COUNT({col}41,{col}43)=2,{col}41>0)'
        put(m,f'{col}44',f'=IF({gate},({col}43-{col}41)/{col}41,"N/D")',PCT)
        put(m,f'{col}45',f'=IF({gate},{col}43-{col}41,"N/D")',USD)
        put(m,f'{col}46',f'=IF(COUNT({col}45)=1,IF({col}45>=0,1,0),"N/D")','0')
    note(m,30,'Lag retrasa el primerfondeo. Fondeos programados después de diciembre no cuentan para outcome ni saldo/valor2027.')
    note(m,31,'Referencia y captura son hipótesis, no baseline post-lanzamiento. La demanda de mercado tampoco equivale a personas.')
    note(m,48,'Finance aplica tasa neta anual /12 al saldo medio mensual externo. Nueva cohorte se pondera por días activa dentro del mes.')
    note(m,49,'La cuenta creada en diciembre no aporta doce meses de valor. Valor2028 no se contabiliza automáticamente.')
    note(m,50,'Revisiones trimestrales optimizan el programa anual; la ejecución y los accesos se comprueban desdeM1.')
    # Cohort schedules, with origin external fixed once and dated monthly cash.
    for c in 'CDEFGHIJK':co.column_dimensions[c].width=15
    put(co,'C4','Escenario')
    head(co,7,['Primerfondeo','Cuentas potenciales','Ticket USD eq.','Origen externo','Ret30d','Ret60d','Ret90d','Ret180d','Ret365d']+[f'M{i+1}' for i in range(12)])
    for i,cc in enumerate(CCOLS):put(co,f'{cc}8',f"='Supuestos'!{COLS[i]}21",'mmm-yy')
    for i,col in enumerate(COLS):
        r=9+i
        put(co,f'C{r}',f'=IF(COUNT(\'Supuestos\'!$D$18,\'Supuestos\'!$D$19)=2,DATE(YEAR(EDATE(\'Supuestos\'!{col}21,\'Supuestos\'!$D$18)),MONTH(EDATE(\'Supuestos\'!{col}21,\'Supuestos\'!$D$18)),\'Supuestos\'!$D$19),"N/D")','dd/mm/yy')
        put(co,f'D{r}',f"='Mensual'!{col}25",NUM);put(co,f'E{r}',"=IF(COUNT('Supuestos'!$D$10)=1,'Supuestos'!$D$10,\"N/D\")",USD)
        for dest,src in (('F',80),('G',55),('H',60),('I',65),('J',70),('K',75)):put(co,f'{dest}{r}',f"='Supuestos'!{col}{src}",PCT)
        for j,cc in enumerate(CCOLS):
            age_r=110+i;ret_r=129+i;ext_r=30+i;avg_r=50+i;avgext_r=70+i
            age=f'{cc}{age_r}';ratio=f'{cc}{ret_r}'
            put(co,age,f'=IF(COUNT($C{r})=1,EOMONTH({cc}$8,0)-$C{r},"N/D")','0')
            curve=f'IF({age}<30,1+($G{r}-1)*{age}/30,IF({age}<60,$G{r}+($H{r}-$G{r})*({age}-30)/30,IF({age}<90,$H{r}+($I{r}-$H{r})*({age}-60)/30,IF({age}<180,$I{r}+($J{r}-$I{r})*({age}-90)/90,IF({age}<=365,$J{r}+($K{r}-$J{r})*({age}-180)/185,$K{r})))))'
            put(co,ratio,f'=IF(COUNT({age},$G{r}:$K{r})=6,IF({age}<0,0,{curve}),"N/D")',PCT)
            put(co,f'{cc}{r}','='+product_guard([f'$D{r}',f'$E{r}',ratio]),USD)
            put(co,f'{cc}{ext_r}','='+product_guard([f'{cc}{r}',f'$F{r}']),USD)
            prior=f'{CCOLS[j-1]}{r}' if j else '0'
            form=f'IF(COUNT($C{r},$D{r},$E{r},{cc}{r},{prior})=5,IF($C{r}>EOMONTH({cc}$8,0),0,IF($C{r}>={cc}$8,($D{r}*$E{r}+{cc}{r})/2*(EOMONTH({cc}$8,0)-$C{r}+1)/DAY(EOMONTH({cc}$8,0)),({prior}+{cc}{r})/2)),"N/D")'
            put(co,f'{cc}{avg_r}','='+form,USD)
            put(co,f'{cc}{avgext_r}','='+product_guard([f'{cc}{avg_r}',f'$F{r}']),USD)
        for labelrow in (30+i,50+i,70+i,110+i,129+i):put(co,f'C{labelrow}',f'=C{r}','dd/mm/yy')
    for top,title in ((28,'Saldo externo al cierre, por cohorte'),(48,'Saldo medio mensual, por cohorte'),(68,'Saldo medio externo mensual, por cohorte'),(108,'Edad de cohorte en días al cierre'),(127,'Saldo relativo interpolado, por cohorte')):
        head(co,top,[title]+['']*8+[f'M{i+1}' for i in range(12)])
        co.merge_cells(start_row=top,start_column=3,end_row=top,end_column=11)
        co[f'C{top}'].alignment=Alignment(horizontal='left',indent=1,vertical='center')
    for row,first,last,label in ((22,9,20,'Saldo cierre USD eq.'),(43,30,41,'Saldo externo cierre'),(63,50,61,'Saldo medio mensual'),(83,70,81,'Saldo medio externo mensual')):
        put(co,f'C{row}',label)
        for cc in CCOLS:put(co,f'{cc}{row}',f'=IF(COUNT({cc}{first}:{cc}{last})=12,SUM({cc}{first}:{cc}{last}),"N/D")',USD)
    head(co,87,['Primerfondeo','Cuentas potenciales','Saldo30d','Saldo60d','Saldo90d','Saldo180d','Saldo365d'])
    for i in range(12):
        src=9+i;r=88+i;put(co,f'C{r}',f'=C{src}','dd/mm/yy');put(co,f'D{r}',f'=D{src}',NUM)
        for dest,days,retcol in (('E',30,'G'),('F',60,'H'),('G',90,'I'),('H',180,'J'),('I',365,'K')):
            form=f'IF(COUNT($C{src})=1,IF($C{src}+{days}>EOMONTH(\'Supuestos\'!$Q$21,0),"Fuera2027",{product_guard([f"$D{src}",f"$E{src}",f"${retcol}{src}"])}),"N/D")'
            put(co,f'{dest}{r}','='+form,USD)
    note(co,102,'365d se muestra Fuera2027 si el aniversario cae en2028. El knot365d sólo ayuda a interpolar la curva hipotética dentro del año.',23)
    note(co,103,'Día15 y lag son supuestos de calendario. Saldo medio usa media de extremos con ponderación de días del primer mes; no dato bancario.',23)
    note(co,104,'La proporción externa de cada fila se conserva en toda su vida. Cambiar el driver de diciembre no repondera cohortes anteriores.',23)
    # Raw public research and historic series, isolated from projections.
    raw.column_dimensions['C'].width=25;raw.column_dimensions['D'].width=54;raw.column_dimensions['E'].width=16;raw.column_dimensions['F'].width=18;raw.column_dimensions['G'].width=16
    head(raw,7,['Grupo','Consulta / territorio','Proxy por mes','Carril','Incluido directo'])
    for i,g in enumerate(source['groups'],8):
        lane={'acquisition_direct':'Directa','acquisition_conditional':'Condicional','assisted_education':'Educación asistida','future_usd':'USD futuro'}.get(g.get('category'),g.get('category',''))
        for col,v in [('C',g['id']),('D',g.get('representative_keyword',g.get('name',''))),('E',g['monthly_volume_proxy']),('F',lane),('G',g['monthly_volume_proxy'] if g.get('use_in_default_envelope') else 0)]:put(raw,f'{col}{i}',v,NUM if col in ('E','G') else None)
    # Total kept in a stable cell separate from raw rows.
    totalrow=len(source['groups'])+9
    put(raw,f'C{totalrow}','Proxy adquisición directa mensual');put(raw,f'G{totalrow}',f'=SUM(G8:G{len(source["groups"])+7})',NUM)
    put(raw,'I5','Proxy directo');put(raw,'J5',f'=G{totalrow}',NUM)
    for r,txt in enumerate(source['metadata']['limitations'],totalrow+3):note(raw,r,txt,11)
    note(raw,max(totalrow+len(source["metadata"]["limitations"])+5,39),'M1 son controles de acceso/medición/delivery. El servicio propuesto dura todo2027; no es piloto de3/6meses.',11)
    head(hist,7,['Observación histórica','Unidad','']+[date(2026,i,1).strftime('%b') for i in range(1,13)])
    put(hist,'C9','Demanda observada por mes calendario');put(hist,'D9','búsquedas')
    put(hist,'C10','Periodo representativo de la muestra');put(hist,'D10',source.get('history_period','Pendiente'))
    put(hist,'C11','Factor / media histórica disponible');put(hist,'D11','factor')
    for i,col in enumerate(COLS):
        hv=source.get('history_monthly_proxy_by_calendar_month',[None]*12)[i]
        put(hist,f'{col}9',hv,NUM)
        put(hist,f'{col}11',f'=IF(COUNT(F9:Q9)=12,IF(AVERAGE(F9:Q9)>0,{col}9/AVERAGE(F9:Q9),"N/D"),"N/D")','0.00')
    note(hist,14,'Historia conserva serie de mercado; no es demanda Pibank post-lanzamiento ni forecast2027. Plano=1 es hipótesis neutral.')
    note(hist,15,'Sólo se habilita Historia cuando los12meses de la misma canasta tienen evidencia. No rellenar meses ausentes con cero.')
    note(hist,16,'La selección de representantes manuales, variantes y cobertura están en source-demand.json. Queda solapamiento entre grupos.')
    # Simple genuine sensitivity calculations.
    head(se,7,['Banda meta provisional','Cuentas requeridas','Visitas requeridas','Depósito inicial USD eq.'])
    for r,p in ((8,.01),(9,.05),(10,.10)):
        put(se,f'C{r}',p,PCT,editable=True);put(se,f'D{r}',f"=C{r}*'Supuestos'!D8",NUM)
        put(se,f'E{r}',f'=IF(COUNT(\'Supuestos\'!F40,\'Supuestos\'!F45)=2,IF(\'Supuestos\'!F40*\'Supuestos\'!F45>0,D{r}/(\'Supuestos\'!F40*\'Supuestos\'!F45),"N/D"),"N/D")',NUM)
        put(se,f'F{r}',f"=D{r}*'Supuestos'!D10",USD)
    se.column_dimensions['E'].width=20
    note(se,12,'Bandas son escala aritmética, no compromiso. Usa tasasM1 constantes; lag y variaciones posteriores se revisan en build anual.',11)
    head(se,15,['Apertura / Fondeo','','',.4,.5,.65,.75,.85])
    for r,p in enumerate([.01,.015,.02,.025,.03],16):
        put(se,f'C{r}',p,PCT,editable=True)
        for col in 'FGHIJ':se[f'{col}15'].number_format=PCT;put(se,f'{col}{r}',f'=IF(COUNT(\'Mensual\'!$F$16:$Q$16)=12,SUM(\'Mensual\'!$F$16:$Q$16)*$C{r}*{col}$15,"N/D")',NUM)
    head(se,23,['Lag a primerfondeo (meses)','Fondeadas dentro2027','Fondeadas después2027'])
    for r,lag in ((24,0),(25,1),(26,2)):
        put(se,f'C{r}',lag,'0',editable=True)
        put(se,f'D{r}',f'=IF(AND(COUNT(\'Mensual\'!$F$25:$Q$25)=12,C{r}>=0,C{r}<12),SUM(\'Mensual\'!$F$25:INDEX(\'Mensual\'!$F$25:$Q$25,1,12-C{r})),"N/D")',NUM)
        put(se,f'E{r}',f'=IF(COUNT(D{r})=1,SUM(\'Mensual\'!$F$25:$Q$25)-D{r},"N/D")',NUM)
    head(se,29,['Inversión hipotética anual USD','Valor/cuenta mínimo para equilibrio','Naturaleza'])
    for r,cost in enumerate(cfg['investment_sensitivity_usd'],30):
        put(se,f'C{r}',cost,USD,editable=True);put(se,f'D{r}',f'=IF(AND(COUNT(\'Resumen\'!$D$10)=1,\'Resumen\'!$D$10>0),C{r}/\'Resumen\'!$D$10,"N/D")',USD);put(se,f'E{r}','Sensibilidad; no fee')
    note(se,34,'Inversión10k/25k/50k es sensibilidad hipotética. No es oferta ni costoEfeonce. ROI sigue N/D hasta inversión y Finance validados.',11)
    head(se,37,['Todas cohortes con origen externo','Saldo externo cierre hipotético'])
    for r,p in ((38,.4),(39,.7),(40,1)):
        put(se,f'C{r}',p,PCT,editable=True);put(se,f'D{r}',f'=IF(COUNT(\'Mensual\'!$Q$35)=1,\'Mensual\'!$Q$35*C{r},"N/D")',USD)
    note(se,43,'Sensibilidad uniforme externa es prueba aparte. El build real conserva proporción de origen de cada cohorte.',11)
    note(se,44,'Retención y ticket se editan en Supuestos y recalculan todos los cierres/valores. Datos ausentes se propaganN/D.',11)
    # Evidence ledger contains scope, authority and no invented measured conversions.
    f.column_dimensions['C'].width=19;f.column_dimensions['D'].width=31;f.column_dimensions['E'].width=62;f.column_dimensions['F'].width=68;f.column_dimensions['G'].width=28
    head(f,7,['ID','Naturaleza','Fuente','Alcance / condición','Estado'])
    entries=[('META','Inferencia provisional','Reunión06/10 + operador','15200×2500=US38m. Año exacto/meta cuenta fondeada no confirmados','Provisional'),
             ('DEMANDA','Estimación de mercado','source-demand.json; DataForSEO','País/idioma, cobertura y periodo en metadata. NoTAM ni usuarios únicos','Observado con límites'),
             ('HISTORIA','Serie histórica de mercado','source-demand.json','No2027 ni Pibank analytics. No crecimiento2x automático','Historia separada'),
             ('SIM','Hipótesis aritméticas','inputs.json','Captura, referencia, apertura, fondeo, relevancia y retención no medidos','Ilustrativo'),
             ('CALENDARIO','Hipótesis de ejecución','inputs.json','Enero2027,12meses,día15+lag. Inicio y fondeo se validan al implementar','Ilustrativo'),
             ('ORIGEN','Hipótesis por cohorte','Supuestos filas80-83','Dinero externo se fija por cohorte; traslado interno no fondeo nuevo grupo','Pendiente banco'),
             ('FINANCE','Autoridad banco','Método neto por definir','Tasa anual neta por USD de saldo medio externo. Un único método, sin dobleFTP','N/D'),
             ('INVERSIÓN','Propuesta no aprobada','Supuestos filas22-24','Inversióncliente vacía. CosteEfeonce vacío. Matriz inversión separada','N/D'),
             ('CAPACIDAD','Estimación planificación','../propuesta-anual/CAPACIDAD-Y-COSTEO-INTERNO.md','2560h+336hreserva=2896h. No fee, roster ni capacidad comprometida','Operations/Finance pendiente'),
             ('GOOGLE','Fuente técnica primaria','https://developers.google.com/search/docs/appearance/ai-features','FundamentosSEO, sin requisitos extra ni garantía; sóloGoogle','Consultada06/10'),
             ('PIBANK','Fuente producto primaria','https://pibank.pe/cuenta-soles-pibank/','CuentaSoles pública. USD eq.modelado no confirma cuentaUSD','Consultada06/10')]
    for r,item in enumerate(entries,8):
        for col,v in zip('CDEFG',item):put(f,f'{col}{r}',v).alignment=Alignment(wrap_text=True,vertical='top')
        f.row_dimensions[r].height=45
    note(f,22,'M1 revisa accesos/datos/delivery; no valida fracaso o éxitoSEO. Programa2027 anual con optimización trimestral.',11)
    note(f,23,'El contrafactual no demuestra pérdidas históricas. CitasIA no se multiplican para fabricar conversiones.',11)
    # Primary output.
    s.column_dimensions['E'].width=17
    note(s,5,'12meses2027. Hipótesis condicionadas para preparar la propuesta; no forecast bancario ni retorno garantizado.',11)
    head(s,7,['Resultado anual dentro2027','Caso activo','Unidad'])
    outputs={8:('Visitas adicionales',"=IF(COUNT('Mensual'!F16:Q16)=12,SUM('Mensual'!F16:Q16),\"N/D\")",'visitas',NUM),
             9:('Aperturas incrementales',"=IF(COUNT('Mensual'!F21:Q21)=12,SUM('Mensual'!F21:Q21),\"N/D\")",'cuentas',NUM),
             10:('Primerfondeo incremental dentro2027',"=IF(COUNT('Mensual'!F28:Q28)=12,SUM('Mensual'!F28:Q28),\"N/D\")",'cuentas',NUM),
             11:('Primerfondeo programado después2027',"=IF(COUNT('Mensual'!F25:Q25,'Mensual'!F28:Q28)=24,SUM('Mensual'!F25:Q25)-SUM('Mensual'!F28:Q28),\"N/D\")",'cuentas',NUM),
             12:('Depósito inicial bruto dentro2027','=IF(COUNT(D10,\'Supuestos\'!D10)=2,D10*\'Supuestos\'!D10,"N/D")','USD eq.',USD),
             13:('Saldo mantenido al cierre2027',"='Mensual'!Q35",'USD eq.',USD),14:('Saldo externo al cierre2027',"='Mensual'!Q36",'USD eq.',USD),
             15:('Saldo medio externo en2027',"=IF(COUNT('Mensual'!F38:Q38)=12,AVERAGE('Mensual'!F38:Q38),\"N/D\")",'USD eq.',USD),
             16:('Escala/meta cuentas provisional','=IF(AND(COUNT(D10,\'Supuestos\'!D8)=2,\'Supuestos\'!D8>0),D10/\'Supuestos\'!D8,"N/D")','%',PCT),
             18:('Inversión cliente anual pendiente',"='Supuestos'!D23",'USD',USD),
             19:('CAC de cuenta fondeada incremental','=IF(AND(\'Supuestos\'!D16="Sí",COUNT(D18,D10)=2,D18>0,D10>0),D18/D10,"N/D")','USD',USD),
             20:('ROI dentro2027',"='Mensual'!Q44",'%',PCT),
             21:('Payback dentro2027','=IF(COUNT(\'Mensual\'!F46:Q46)=12,IFERROR(MATCH(1,\'Mensual\'!F46:Q46,0),"No en2027"),"N/D")','mes',NUM)}
    for r,(label,formula,unit,fmt) in outputs.items():
        put(s,f'C{r}',label);put(s,f'D{r}',formula,RATE if r==16 else fmt).font=Font(name='Arial',size=10,color='18294A');put(s,f'E{r}',unit).alignment=Alignment(indent=2,vertical='center')
    note(s,23,'Saldo no es revenue. ROI/payback requieren inversión y Finance validados. Ninguna cohorte aporta valor después del horizonte.',11)
    head(s,25,['Mes','Visitas adicionales','Primerfondeo inc.','Saldo cierre USD eq.','Saldo externo USD eq.','Saldo medio externo'])
    for i,col in enumerate(COLS):
        r=26+i
        for dest,src in (('C',8),('D',16),('E',28),('F',35),('G',36),('H',38)):put(s,f'{dest}{r}',f"='Mensual'!{col}{src}",'mmm-yy' if dest=='C' else USD if dest in 'FGH' else NUM).font=Font(name='Arial',size=10,color='18294A')
    note(s,40,'M1: acceso/medición/delivery. M3/6/9/12: optimización trimestral. No pilotoSEO ni juicio de viabilidad a3meses.',11)
    note(s,41,'Demanda histórica y alcance de investigación se leen en Demanda/Historia; los drivers2027 continúan siendo hipótesis.',11)
    note(s,42,'US$38m y ticketUS$2500 provisionales; anualidad de la meta aún no confirmada. USD eq.no significa productoUSD disponible.',11)
    note(s,43,'Origen externo fijo por cohorte. Saldo medio pondera días activos; diciembre tiene menos tiempo de contribución que enero.',11)
    note(s,45,'Cambiar selector en SupuestosD4. Cambiar lagD18/díaD19 o tasas actualiza el mismo build; fuentes históricas permanecen intactas.',11)
    wb.save(output)

def recalculate(path,out):
    out.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='pibank-annual-lo-') as profile:
        p=subprocess.run([str(SOFFICE),f'-env:UserInstallation={Path(profile).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=90)
        if p.returncode or not (out/path.name).exists():raise RuntimeError(f'{p.returncode}: {p.stdout} {p.stderr}')
    return out/path.name

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--no-copy',action='store_true');args=parser.parse_args()
    cfg=json.loads((HERE/'inputs.json').read_text());source=json.loads((HERE/'source-demand.json').read_text())
    assert cfg['months']==12 and cfg['selected_case'] in (1,2,3)
    assert isinstance(cfg['demand_monthly_proxy'],(int,float)) and cfg['demand_monthly_proxy']>=0
    assert cfg['demand_monthly_proxy']==source['metadata']['default_envelope_searches_monthly']
    assert sum(cfg['monthly_base_capacity_hours'])==cfg['delivery_capacity_hours_base']
    assert cfg['delivery_capacity_hours_base']+cfg['delivery_capacity_hours_reserve']==cfg['delivery_capacity_hours_total']
    for cs in cfg['cases']:
        assert len(cs['extra_capture'])==12 and all(0<=v+cs['reference_capture']<=1 for v in cs['extra_capture'])
        knots=[cs[k] for k in ('retention30','retention60','retention90','retention180','retention365')]
        assert all(0<=x<=1 for x in knots) and knots==sorted(knots,reverse=True)
    with tempfile.TemporaryDirectory(prefix='pibank-annual-build-') as tmp:
        raw=Path(tmp)/'raw';raw.mkdir();file=raw/NAME;build(cfg,source,file)
        final=recalculate(file,Path(tmp)/'calc');stage=HERE/f'.{NAME}.building';shutil.copy2(final,stage);stage.replace(HERE/NAME)
    results={'as_of':cfg['as_of'],'horizon':'2027,12meses desde enero ilustrativo; no incluye2028','goal_status':cfg['goal_status'],'demand_monthly_proxy':cfg['demand_monthly_proxy'],
             'demand_monthly':cfg['demand_monthly_proxy'],'demand_status':cfg['demand_status'],'relevance_factor_illustrative':cfg['relevance_eligibility_factor'],'growth_multiplier_hypothesis':cfg['growth_multiplier'],
             'seasonality_mode':cfg['seasonality_mode'],'cases':[simulate(cfg,i) for i in (1,2,3)],'workbook':NAME,'recalculation_engine':'Bundled LibreOffice'}
    for cs in results['cases']:
        cs['funded_incremental']=cs['funded_incremental_within2027'];cs['account_goal_share']=cs['account_goal_share_arithmetic']
    base=cfg['cases'][1];out=results['cases'][1]
    results['scale_tests_base']=[{'share_provisional_goal':p,'accounts':cfg['goal_accounts']*p,'visits_required_at_constant_base_rates':cfg['goal_accounts']*p/(base['opening_rate']*base['funding_rate']) if base['opening_rate']*base['funding_rate']>0 else None,
                                 'nature':'Escala hipotética; no cuota prometida ni periodo confirmado meta'} for p in (.01,.05,.10)]
    results['investment_sensitivity_base']=[{'hypothetical_total_investment_usd':v,'value_per_account_needed_for_equilibrium':v/out['funded_incremental_within2027'] if out['funded_incremental_within2027']>0 else None,'nature':'Sensibilidad, no fee ni ROI'} for v in cfg['investment_sensitivity_usd']]
    results['delivery_capacity']={'base_hours':cfg['delivery_capacity_hours_base'],'reserve_hours':cfg['delivery_capacity_hours_reserve'],'total_hours':cfg['delivery_capacity_hours_total'],'status':'Estimación planificación, no fee ni costo real'}
    (HERE/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2)+'\n')
    if not args.no_copy:
        DELIVERY.mkdir(parents=True,exist_ok=True)
        for p in HERE.iterdir():
            if p.is_file() and p.suffix in ('.json','.md','.py','.xlsx'):shutil.copy2(p,DELIVERY/p.name)
    print(json.dumps({'workbook':str(HERE/NAME),'outputs':[{k:v for k,v in cs.items() if k!='monthly'} for cs in results['cases']]},ensure_ascii=False))

if __name__=='__main__':main()
