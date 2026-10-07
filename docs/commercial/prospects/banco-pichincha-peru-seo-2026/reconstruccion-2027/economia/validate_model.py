"""Independent reverse-model and native control tests; generator is not imported."""
from pathlib import Path
from datetime import date,datetime
import calendar,csv,hashlib,json,math,subprocess,tempfile
import openpyxl
HERE=Path(__file__).resolve().parent;NAME='Pibank-decision-inversion-2027.xlsx';SOFFICE='/Users/jreye/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice'
COLS=[openpyxl.utils.get_column_letter(i)for i in range(6,18)];CCOLS=[openpyxl.utils.get_column_letter(i)for i in range(12,24)]
def close(x,y,label):assert isinstance(x,(int,float))and math.isclose(x,y,abs_tol=1e-6,rel_tol=1e-9),(label,x,y)
def relative(age,ys):
    if age<0:return 0
    xs=[0,30,60,90,180,365];ys=[1]+ys
    for i in range(1,6):
        if age<=xs[i]:return ys[i-1]+(ys[i]-ys[i-1])*(age-xs[i-1])/(xs[i]-xs[i-1])
    return ys[-1]
def independent(w):
    for ws in w:
        for row in ws:
            for c in row:assert c.data_type!='e',(ws.title,c.coordinate,c.value)
    a=w['Supuestos'];r=w['Requisitos'];co=w['Cohortes'];d=w['Decision'];e=w['Economia'];target=a['D12'].value;typ=a['G6'].value;window=a['D45'].value
    if not isinstance(a['D9'].value,(int,float))or a['D9'].value==0:
        assert d['D9'].value=='N/D'
        return {'accounts_required':target,'definition':typ,'visits_required':'N/D','funded_within_window':d['D20'].value,'status':'Unknown/infeasible visit requirement; cohort quality remains conditional on the required accounts'}
    critical=[a[c].value for c in('D10','D11','D13','D14','D45')]+[a[f'{c}34'].value for c in COLS]
    if not all(isinstance(x,(int,float))for x in critical)or not isinstance(a['D36'].value,datetime)or (typ=='Fondeadas'and a['D10'].value==0):
        assert d['D11'].value=='N/D' and e['D8'].value=='N/D'
        return {'accounts_required':target,'definition':typ,'visits_required':d['D9'].value,'funded_within_window':d['D20'].value,'status':'N/D from missing/infeasible requirement input'}
    start=a['D36'].value.date();lag=a['D13'].value;fundrate=a['D10'].value;openingrate=a['D9'].value;weights=[a[f'{c}34'].value if i<window and(typ=='Aperturas'or i+lag<window)else 0 for i,c in enumerate(COLS)]
    if sum(weights)==0:
        assert d['D9'].value=='N/D' and d['D11'].value=='N/D';return {'accounts_required':target,'definition':typ,'status':'N/D, no eligible weights'}
    opened=[target*(1 if typ=='Aperturas'else 1/fundrate)*v/sum(weights)for v in weights]
    funded=[opened[i-lag]*fundrate if i<window and i>=lag else 0 for i in range(12)];visits=[v/openingrate for v in opened]
    close(d['D9'].value,sum(visits),'requirements');close(d['D20'].value,sum(funded),'funding quality');close(d['D22'].value,sum(opened)*fundrate-sum(funded),'scheduled outside')
    dates=[a[f'{c}31'].value.date()for c in COLS];ends=[v.replace(day=calendar.monthrange(v.year,v.month)[1])for v in dates]
    ratios=[a[f'D{j}'].value for j in(37,38,39,40,41)];stocks=[0.]*12;external=[0.]*12;avgext=[0.]*12
    for i,c in enumerate(COLS):
        fd=max(start,date(dates[i].year,dates[i].month,a['D14'].value));initial=funded[i]*a['D8'].value;origin=a[f'{c}44'].value;prior=0
        assert co[f'C{9+i}'].value.date()==fd
        for j,(cc,dt,end)in enumerate(zip(CCOLS,dates,ends)):
            balance=initial*relative((end-fd).days,ratios)if j<window else 0
            if j>=window or fd>end:mean=0
            elif fd>=dt.replace(day=1):mean=(initial+balance)/2*((end-fd).days+1)/calendar.monthrange(dt.year,dt.month)[1]
            else:mean=(prior+balance)/2
            close(co[f'{cc}{9+i}'].value,balance,'cohort close');close(co[f'{cc}{30+i}'].value,balance*origin,'sticky external origin');close(co[f'{cc}{50+i}'].value,mean*origin,'active days average')
            stocks[j]+=balance;external[j]+=balance*origin;avgext[j]+=mean*origin;prior=balance
        for dest,days,k in(('E',30,0),('F',60,1),('G',90,2),('H',180,3),('I',365,4)):
            actual=co[f'{dest}{67+i}'].value
            if(fd-ends[window-1]).days+days>0:assert actual=='Fuera ventana'
            else:close(actual,initial*ratios[k],'maturity')
    for j,c in enumerate(COLS):close(r[f'{c}19'].value,avgext[j],'average monthly');close(r[f'{c}20'].value,external[j],'end external')
    close(d['D11'].value,sum(avgext)/window,'average within window');close(d['D12'].value,external[window-1],'end within window')
    method=a['D52'].value;rate=a['D17'].value;finance=a['D16'].value=='Sí'
    bankvalue=None
    if finance and method=='Saldo medio externo'and isinstance(rate,(int,float)):bankvalue=sum(avgext)*rate/12
    if finance and method=='Valor neto por fondeada'and isinstance(a['D54'].value,(int,float)):bankvalue=sum(funded)*a['D54'].value
    if bankvalue is None:assert e['D8'].value=='N/D' and e['D9'].value=='N/D'
    else:close(e['D8'].value,bankvalue,'unique net method');close(e['D9'].value,bankvalue/(1+a['D18'].value),'investment ceiling')
    costs=[a[f'{c}32'].value for c in COLS[:window]];ready=finance and method=='Saldo medio externo'and isinstance(rate,(int,float))and a['D19'].value=='Sí'and all(isinstance(x,(int,float))and x>=0 for x in costs)and sum(costs)>0
    if ready:
        close(d['D18'].value,(bankvalue-sum(costs))/sum(costs),'conditional ROI within window')
        payback=next((j+1 for j in range(window)if sum(avgext[:j+1])*rate/12>=sum(costs[:j+1])),'No en ventana')
        assert d['D19'].value==payback
    else:assert d['D18'].value=='N/D'and d['D19'].value=='N/D'
    loaded=a['D21'].value;others=a['D25'].value;margin=a['D22'].value;internal=a['D20'].value=='Sí'and isinstance(loaded,(int,float))and loaded>=0 and isinstance(others,(int,float))and others>=0 and isinstance(margin,(int,float))and 0<=margin<1
    if internal:
        cost=a['D23'].value*loaded+others;close(e['D13'].value,cost,'internal costs');close(e['D15'].value,cost/(1-margin),'cost floor')
        if bankvalue is not None and loaded>0:close(e['D17'].value,max(0,(bankvalue/(1+a['D18'].value)*(1-margin)-others)/loaded),'capacity compatibility')
    else:assert e['D15'].value=='N/D'
    for row,share in((25,.01),(26,.05),(27,.10)):
        close(d[f'D{row}'].value,a['D6'].value*share,'band requirement');close(d[f'E{row}'].value,sum(visits)*a['D6'].value*share/target,'band visits')
    for row in(24,25,26,27):close(e[f'D{row}'].value,sum(avgext)*e[f'C{row}'].value/12,'net value sensitivity')
    return {'accounts_required':target,'definition':typ,'visits_required':sum(visits),'funded_within_window':sum(funded),'funded_after_window':sum(opened)*fundrate-sum(funded),'average_external_balance':sum(avgext)/window,'external_balance_end':external[window-1],'bank_value':e['D8'].value,'investment_ceiling':e['D9'].value,'cost_floor':e['D15'].value,'roi_within_window':d['D18'].value,'status':'checked independent within window'}

def native(path,out):
    out.mkdir()
    with tempfile.TemporaryDirectory(prefix='pibank-inverse-test-profile-')as profile:
        p=subprocess.run([SOFFICE,f'-env:UserInstallation={Path(profile).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=90)
        assert p.returncode==0 and(out/path.name).exists(),(p.stdout,p.stderr)
    return openpyxl.load_workbook(out/path.name,data_only=True)

def main():
    path=HERE/NAME;checks=[{'test':'Delivered cache: annual Aperturas default, cohorts and requirements','result':independent(openpyxl.load_workbook(path,data_only=True))}]
    capacity=list(csv.DictReader((HERE.parent/'capacidad-programa-con-pr.csv').open()))
    cached=openpyxl.load_workbook(path,data_only=True)
    assert sum(int(r['total_h'])for r in capacity)==cached['Supuestos']['D23'].value==cached['CosteoRoles']['D22'].value==1216
    assert sum(int(r['reserva_h'])for r in capacity)==cached['CosteoRoles']['E22'].value==150
    assert cached['Economia']['F33'].value==1216
    for i,r in enumerate(capacity,8):
        assert cached['CosteoRoles'][f'C{i}'].value==r['rol']
        assert cached['CosteoRoles'][f'D{i}'].value==int(r['total_h'])
    calendar_rows=list(csv.DictReader((HERE.parent/'CALENDARIO-INTEGRADO-2027.csv').open()))
    assert len(calendar_rows)==12
    for i,r in enumerate(calendar_rows,8):assert cached['Plan2027'][f'D{i}'].value==r['entrega_principal']
    assert cached['CosteoRoles']['G22'].value=='N/D' and cached['Supuestos']['D21'].value=='N/D'
    financed={'D16':'Sí','D17':.02,'D19':'Sí',**{f'{c}32':1000 for c in COLS}};costed={**financed,'D20':'Sí','D21':30,'D22':.45,'D25':0}
    mods=[('5pct opening band',{'D3':2}),('10pct opening band',{'D3':3}),('Funded target explicitly selected',{'G6':'Fondeadas'}),('Funding lag1 opening target',{'D13':1}),('Funding lag1 funded target',{'G6':'Fondeadas','D13':1}),('Funding lag12 all funding outside',{'D13':12}),('Last origin external changed only',{'Q44':.3}),('Finance alone no investment',{'D16':'Sí','D17':.02}),('Finance+investment within2027',financed),('Incomplete investment',{**financed,'Q32':None}),('Zero investment distinct from blank',{**financed,**{f'{c}32':0 for c in COLS}}),('Negative expense refused',{**financed,'Q32':-10}),('Zero expense in one month valid',{**financed,'Q32':0}),('Costed floor/capacity compatibility',costed),('Unknown other internal cost',{**costed,'D25':None}),('Zero loaded cost not unlimited capacity',{**costed,'D21':0}),('Retrolaunch window unknown',{'D29':'Retrolanzamiento'}),('Retrolaunch3months from partial launch month',{**financed,'D29':'Retrolanzamiento','D30':datetime(2026,9,14),'D35':3}),('Retrolaunch zero months invalid',{'D29':'Retrolanzamiento','D30':datetime(2026,9,14),'D35':0}),('24m approved valueperaccount method no double count',{**financed,'D52':'Valor neto por fondeada','D53':24,'D54':50}),('36m unknown peraccount remains ND',{'D16':'Sí','D52':'Valor neto por fondeada','D53':36}),('Zero funding opening requirement survives',{'D10':0}),('Zero funding funded requirement infeasible',{'G6':'Fondeadas','D10':0}),('Missing selected opening rate',{'D9':None})]
    rolecost={f'CosteoRoles!F{i}':30 for i in range(8,21)}
    rolevalidated={**rolecost,'D20':'Sí','D22':.45,'D25':0}
    mods.extend([('All13roles fully loaded cost',rolevalidated),('One missing role refuses blended',{**rolevalidated,'CosteoRoles!F20':None}),('Negative role refuses blended',{**rolevalidated,'CosteoRoles!F20':-1}),('Explicit zero role permitted',{**rolevalidated,'CosteoRoles!F20':0})])
    with tempfile.TemporaryDirectory(prefix='pibank-inverse-validation-')as tmp:
        for i,(label,changes)in enumerate(mods):
            w=openpyxl.load_workbook(path)
            for cell,value in changes.items():
                sh,addr=cell.split('!',1)if '!'in cell else('Supuestos',cell)
                w[sh][addr]=value
            p=Path(tmp)/f'test{i}.xlsx';w.save(p);checked=native(p,Path(tmp)/f'calc{i}');checks.append({'test':label,'changed_inputs':{k:v.isoformat()if isinstance(v,datetime)else v for k,v in changes.items()},'result':independent(checked)})
            if label.startswith('All13roles'):
                close(checked['CosteoRoles']['G22'].value,1216*30,'13 roles cost sum')
                close(checked['Supuestos']['D21'].value,30,'weighted blended')
            if label.startswith(('One missing role','Negative role')):assert checked['CosteoRoles']['G22'].value=='N/D' and checked['Supuestos']['D21'].value=='N/D'
            if label.startswith('Explicit zero role'):close(checked['CosteoRoles']['G22'].value,(1216-int(capacity[-1]['total_h']))*30,'explicit zero role')
    sources=json.loads((HERE/'sources.json').read_text())
    for entry in sources['historical_artifacts']:assert hashlib.sha256(Path(entry['path']).read_bytes()).hexdigest()==entry['sha256'],'Preserved historical artifact changed'
    formulas=openpyxl.load_workbook(path,data_only=False);assert formulas['Decision']['D9'].data_type=='f'and formulas['Cohortes']['W20'].data_type=='f';assert formulas['Supuestos']['D29'].value=='Programa2027'
    result={'status':'passed','native_recalculation_tests':len(checks),'engine':'Bundled LibreOffice','workbook_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'checks':checks,'combined_capacity_calendar_checked':True,'historical_artifacts_unchanged':True,'unverified':'Excel Desktop UI, bank value, quote and Operations capacity approval.'}
    (HERE/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');print(json.dumps({'status':'passed','tests':len(checks),'default':checks[0]['result']},ensure_ascii=False))
if __name__=='__main__':main()
