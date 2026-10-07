"""Decision-first reverse investment model; not a bank forecast or approved quote."""
from pathlib import Path
from datetime import date
import argparse, calendar, hashlib, json, math, shutil, subprocess, tempfile
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

HERE=Path(__file__).resolve().parent
NAME='Pibank-decision-inversion-2027.xlsx'
OUT=Path('/Users/jreye/Documents/Banco Pichincha Peru — Prospect Case/02-Uso-interno/reconstruccion-2027/economia')
SOFFICE='/Users/jreye/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice'
COLS=[openpyxl.utils.get_column_letter(i) for i in range(6,18)]
CCOLS=[openpyxl.utils.get_column_letter(i) for i in range(12,24)]
NUM='#,##0.0;(#,##0.0);"—"';MONEY='"US$"#,##0;("US$"#,##0);"—"';PCT='0.00%;(0.00%);"—"'

def put(ws,cell,value,fmt=None,input=False):
    c=ws[cell];c.value=value;c.font=Font(name='Arial',size=10,color='0000CC' if input else '008050' if isinstance(value,str) and value.startswith('=') and '!' in value else '18294A')
    c.alignment=Alignment(vertical='center',horizontal='right' if fmt else 'left',indent=1 if fmt else 0)
    if fmt:c.number_format=fmt
    if input:c.fill=PatternFill('solid',fgColor='FFF2CC')
    return c

def note(ws,row,value,last=17):
    ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=last);put(ws,f'C{row}',value).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row].height=32

def header(ws,row,labels,start=3):
    for i,v in enumerate(labels,start):
        c=put(ws,f'{openpyxl.utils.get_column_letter(i)}{row}',v);c.fill=PatternFill('solid',fgColor='20385F');c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
    ws.row_dimensions[row].height=30

def sheet(wb,name,title,last=45,end=17):
    w=wb.create_sheet(name);w.sheet_view.showGridLines=False;w.freeze_panes='F8';w.column_dimensions['A'].width=w.column_dimensions['B'].width=2;w.column_dimensions['C'].width=43;w.column_dimensions['D'].width=18;w.column_dimensions['E'].width=3
    for i in range(6,end+1):w.column_dimensions[openpyxl.utils.get_column_letter(i)].width=14
    for r in range(1,last+1):w.row_dimensions[r].height=20
    put(w,'C2',title).font=Font(name='Arial',size=14,bold=True,color='18294A');w.row_dimensions[2].height=25
    w.page_setup.orientation='landscape';w.page_setup.paperSize=w.PAPERSIZE_A3;w.page_setup.fitToWidth=1;w.page_setup.fitToHeight=1;w.sheet_properties.pageSetUpPr.fitToPage=True;w.print_area=f'C1:{openpyxl.utils.get_column_letter(end)}{last}';w.page_margins.left=w.page_margins.right=.2;w.page_margins.top=w.page_margins.bottom=.25
    return w

def guard(parts,expression):return f'IF(COUNT({",".join(parts)})={len(parts)},{expression},"N/D")'

def build(cfg,source,path):
    wb=openpyxl.Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    d=sheet(wb,'Decision','Pibank 2027: decisión de inversión, no forecast',48,11)
    a=sheet(wb,'Supuestos','Inputs: metas contextuales, pruebas y validaciones',54)
    r=sheet(wb,'Requisitos','Requisitos mensuales para la contribución elegida',40)
    co=sheet(wb,'Cohortes','Saldo externo y valor dentro del horizonte 2027',108,23);co.page_setup.fitToHeight=3
    e=sheet(wb,'Economia','Techo bancario y piso operativo: decisión condicionada',48,11)
    e.column_dimensions['E'].width=22;e.column_dimensions['F'].width=34
    f=sheet(wb,'Evidencia','Cobertura, carriles y autoridad de cada dato',49,11)
    retro=sheet(wb,'Retrolanzamiento','Retrolanzamiento: ventana desconocida y contrafactual',36,11)
    note(a,4,'Bandas 1/5/10% son decisiones de contribución contextual. No objetivos acordados ni previsiones del canal.')
    rows={6:('Meta cuentas provisional',cfg['goal_accounts'],NUM),7:('Meta USD inferida, sin periodo confirmado',cfg['goal_usd'],MONEY),8:('Ticket USD equivalente provisional',cfg['ticket_usd_equivalent'],MONEY),9:('Visita→apertura, prueba aritmética',cfg['opening_rate_test'],PCT),10:('Apertura→fondeo, prueba aritmética',cfg['funding_rate_test'],PCT),11:('Elegibilidad del panel, prueba',cfg['eligibility_test'],PCT),13:('Lag apertura→primer fondeo (meses)',cfg['funding_lag_months'],'0'),14:('Día de primer fondeo (1–28)',cfg['funding_day'],'0'),15:('Inicio ilustrativo programa',date.fromisoformat(cfg['start_month']),'mmm-yy'),16:('Finance valida método/valor neto','Sí' if cfg['finance_validated'] else 'No',None),17:('Valor neto anual por USD saldo medio externo',cfg['net_annual_value_per_external_usd'],PCT),18:('ROI mínimo exigido, prueba 0=equilibrio',cfg['roi_hurdle_test'],PCT),19:('Inversión cliente completa validada','Sí' if cfg['investment_validated'] else 'No',None),20:('Costeo/margen Efeonce validado','Sí' if cfg['internal_cost_validated'] else 'No',None),21:('Costo real blended USD/h (pendiente)',cfg['loaded_cost_usd_hour'],MONEY),22:('Margen bruto autorizado (pendiente)',cfg['approved_margin'],PCT),23:('Horas candidato enfocado (base+reserva)',cfg['focused_capacity_hours_total'],NUM),25:('Otros costos enfocado USD (pendiente)',cfg['other_internal_cost_usd'],MONEY)}
    put(a,'C3','Selector banda (1 / 2 / 3)');put(a,'D3',cfg['selected_band'],input=True)
    put(a,'C5','Participación contextual elegida');put(a,'D5','=CHOOSE(D3,1%,5%,10%)',PCT)
    put(a,'C12','Cuentas fondeadas requeridas dentro de 2027');put(a,'D12','=D5*D6',NUM)
    a['C12']='Cuentas requeridas según definición elegida'
    for row,(label,value,fmt) in rows.items():put(a,f'C{row}',label);put(a,f'D{row}',value,fmt,input=row!=23)
    note(a,27,'Conversión, elegibilidad, ticket, retención, origen y distribución son hipótesis. Se calculan requisitos; no se predicen resultados.')
    note(a,28,'US$38m = 15.200×2.500: inferencia provisional. Programa anual no confirma definición/anualidad de meta ni producto USD.')
    put(a,'F6','Definición de la banda');put(a,'G6',cfg['goal_type'],input=True)
    put(a,'C29','Contexto del análisis');put(a,'D29',cfg['analysis_context'],input=True)
    put(a,'C30','Fecha retrolanzamiento (desconocida)');put(a,'D30',None if cfg['retro_start'] is None else date.fromisoformat(cfg['retro_start']),'dd/mm/yy',input=True)
    put(a,'C35','Meses ventana retrospectiva (desconocidos)');put(a,'D35',cfg['retro_window_months'],'0',input=True)
    put(a,'C36','Inicio del cálculo activo');put(a,'D36','=IF(D29="Programa2027",D15,IF(COUNT(D30)=1,D30,"N/D"))','dd/mm/yy')
    put(a,'C45','Meses activos del cálculo');put(a,'D45','=IF(D29="Programa2027",12,IF(COUNT(D35)=1,IF(AND(D35>=1,D35<=12),D35,"N/D"),"N/D"))','0')
    put(a,'C52','Método único de valor bancario');put(a,'D52',cfg['valuation_method'],input=True)
    put(a,'C53','Horizonte del valor, meses');put(a,'D53',cfg['valuation_horizon_months'],'0',input=True)
    put(a,'C54','Valor neto por fondeada aprobado (ND)');put(a,'D54',cfg['bank_approved_net_value_per_funded_account'],MONEY,input=True)
    for i,c in enumerate(COLS):put(a,f'{c}31',f'=IF(COUNT($D$36)=1,EDATE($D$36,{i}),"N/D")','mmm-yy')
    for row,label,key in ((32,'Inversión cliente USD mensual, pendiente','monthly_client_investment_usd'),(34,'Peso del objetivo de fondeo, prueba','funding_weights'),(44,'Origen externo por cohorte, prueba','external_share_by_cohort')):
        put(a,f'C{row}',label)
        for i,c in enumerate(COLS):put(a,f'{c}{row}',cfg[key][i],PCT if row==44 else MONEY if row==32 else NUM,input=True)
    for row,key,label in ((37,'retention30','Saldo relativo 30 días'),(38,'retention60','Saldo relativo 60 días'),(39,'retention90','Saldo relativo 90 días'),(40,'retention180','Saldo relativo 180 días'),(41,'retention365','Saldo relativo 365 días')):put(a,f'C{row}',label+' (hipótesis)');put(a,f'D{row}',cfg[key],PCT,input=True)
    note(a,46,'Los pesos distribuyen una meta requerida, no una rampa SEO. Lag excluye fondeos que necesitarían aperturas anteriores al programa.')
    note(a,47,'Origen externo se fija por cohorte. No sumar traslado interno como nuevo fondeo del grupo.')
    note(a,48,'M1: accesos/datos/delivery. Revisiones M3/6/9/12: optimización del programa anual, sin piloto ni gate de ROI temprano.')
    note(a,49,'Costo/h incluye sólo el método fully loaded validado; evitar duplicar overhead. Ningún fee o roster está aprobado.')
    for row in (16,19,20):
        v=DataValidation(type='list',formula1='"Sí,No"');a.add_data_validation(v);v.add(a[f'D{row}'])
    for cell,choices in (('G6','Aperturas,Fondeadas'),('D29','Programa2027,Retrolanzamiento'),('D52','Saldo medio externo,Valor neto por fondeada'),('D53','12,24,36')):
        v=DataValidation(type='list',formula1='"'+choices+'"');a.add_data_validation(v);v.add(a[cell])
    for cell,low,high in (('D3',1,3),('D13',0,12),('D14',1,28),('D35',1,12)):
        v=DataValidation(type='whole',operator='between',formula1=low,formula2=high);v.showErrorMessage=True;v.errorStyle='stop';a.add_data_validation(v);v.add(a[cell])
    for rng in ('D9:D11','D22','D37:D41','F44:Q44'):
        v=DataValidation(type='decimal',operator='between',formula1=0,formula2=1,allow_blank=True);v.showErrorMessage=True;v.errorStyle='stop';a.add_data_validation(v);v.add(rng)
    for rng in ('F32:Q32','F34:Q34','D21','D25'):
        v=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0,allow_blank=True);v.showErrorMessage=True;v.errorStyle='stop';a.add_data_validation(v);v.add(rng)
    # Monthly requirement build. Delay reduces funding months rather than assigning pre-program outcomes.
    header(r,6,['Requisito','Unidad','']+[f'M{i+1}'for i in range(12)])
    labels={9:('Peso de fondeo indicado','peso'),10:('Peso elegible por inicio/lag','peso'),11:('Primeras cuentas fondeadas requeridas','cuentas'),12:('Aperturas necesarias desde este mes','cuentas'),13:('Visitas necesarias desde este mes','visitas'),15:('Panel directo observado, contexto parcial','búsquedas/mes'),16:('Panel tras elegibilidad hipotética','búsquedas/mes'),17:('Visitas requeridas / panel elegible','ratio panel'),19:('Saldo externo medio requerido','USD eq.'),20:('Saldo externo al cierre requerido','USD eq.'),22:('Valor neto bancario condicionado','USD'),23:('Inversión cliente pendiente','USD'),24:('Valor acumulado condicionado','USD'),25:('Inversión acumulada','USD'),26:('ROI condicional a logro de la banda','%'),27:('Recuperación alcanzada (condicional)','0/1')}
    for row,(label,unit)in labels.items():put(r,f'C{row}',label);put(r,f'D{row}',unit)
    for i,c in enumerate(COLS):
        put(r,f'{c}7',f"='Supuestos'!{c}31",'mmm-yy');put(r,f'{c}9',f'=IF(COUNT(\'Supuestos\'!{c}34)=1,\'Supuestos\'!{c}34,"N/D")',NUM)
        put(r,f'{c}10',f'=IF(COUNT(\'Supuestos\'!$D$13,\'Supuestos\'!$D$45,{c}9)=3,IF({i}<\'Supuestos\'!$D$45,IF(\'Supuestos\'!$G$6="Fondeadas",IF({i}+\'Supuestos\'!$D$13<\'Supuestos\'!$D$45,{c}9,0),{c}9),0),"N/D")',NUM)
        put(r,f'{c}12',f'=IF(COUNT($F$10:$Q$10)=12,IF(SUM($F$10:$Q$10)>0,IF(\'Supuestos\'!$G$6="Aperturas",\'Supuestos\'!$D$12*{c}10/SUM($F$10:$Q$10),IF(COUNT(\'Supuestos\'!$D$10)=1,IF(\'Supuestos\'!$D$10>0,\'Supuestos\'!$D$12*{c}10/SUM($F$10:$Q$10)/\'Supuestos\'!$D$10,"N/D"),"N/D")),"N/D"),"N/D")',NUM)
        put(r,f'{c}11',f'=IF(COUNT(\'Supuestos\'!$D$13,\'Supuestos\'!$D$45,\'Supuestos\'!$D$10)=3,IF(AND({i}<\'Supuestos\'!$D$45,{i}>=\'Supuestos\'!$D$13),IF(COUNT(INDEX($F$12:$Q$12,1,{i+1}-\'Supuestos\'!$D$13))=1,INDEX($F$12:$Q$12,1,{i+1}-\'Supuestos\'!$D$13)*\'Supuestos\'!$D$10,"N/D"),0),"N/D")',NUM)
        put(r,f'{c}13',f'=IF(COUNT(\'Supuestos\'!$D$9,{c}12)=2,IF(\'Supuestos\'!$D$9>0,{c}12/\'Supuestos\'!$D$9,"N/D"),"N/D")',NUM)
        put(r,f'{c}15',"='Evidencia'!J5",NUM)
        put(r,f'{c}16','='+guard([f'{c}15',"'Supuestos'!$D$11","'Supuestos'!$D$45"],f'IF({i}<\'Supuestos\'!$D$45,{c}15*\'Supuestos\'!$D$11,0)'),NUM)
        put(r,f'{c}17',f'=IF(COUNT({c}13,{c}16)=2,IF({c}16>0,{c}13/{c}16,"N/D"),"N/D")',PCT)
        for row,src in ((19,63),(20,43)):put(r,f'{c}{row}',f"='Cohortes'!{CCOLS[i]}{src}",MONEY)
        put(r,f'{c}22',f'=IF(COUNT(\'Supuestos\'!$D$45)=1,IF({i}>=\'Supuestos\'!$D$45,0,IF(AND(\'Supuestos\'!$D$16="Sí",\'Supuestos\'!$D$52="Saldo medio externo",COUNT(\'Supuestos\'!$D$17,{c}19)=2),{c}19*\'Supuestos\'!$D$17/12,"N/D")),"N/D")',MONEY)
        put(r,f'{c}23',f'=IF(COUNT(\'Supuestos\'!$D$45)=1,IF({i}>=\'Supuestos\'!$D$45,0,IF(COUNT(\'Supuestos\'!{c}32)=1,IF(\'Supuestos\'!{c}32>=0,\'Supuestos\'!{c}32,"N/D"),"N/D")),"N/D")',MONEY)
        for row,base in ((24,22),(25,23)):put(r,f'{c}{row}',f'=IF(COUNT($F${base}:{c}{base})={i+1},SUM($F${base}:{c}{base}),"N/D")',MONEY)
        ready=f'AND(\'Supuestos\'!$D$16="Sí",\'Supuestos\'!$D$19="Sí",COUNT({c}24,{c}25)=2,{c}25>0)'
        put(r,f'{c}26',f'=IF({ready},({c}24-{c}25)/{c}25,"N/D")',PCT)
        put(r,f'{c}27',f'=IF(COUNT({c}26)=1,IF({c}26>=0,1,0),"N/D")','0')
    note(r,30,'Un ratio >100% significa que este panel no alcanza para el requisito bajo esos supuestos; no demuestra inviabilidad del mercado.')
    note(r,31,'El denominador es una canasta parcial de consultas, con overlap residual. No son personas únicas, TAM ni visitas Pibank.')
    note(r,32,'Marca, educación, activación y retención tienen recorridos distintos; sus retornos permanecen N/D sin baseline y deduplicación.')
    note(r,33,'Toda valoración/ROI es condicional al logro de la banda. No transforma requisitos en forecast ni atribuye todo al SEO/AEO.')
    # External average balance is dated by first funding; origin fixed per cohort.
    for c in 'CDEFGHIJK':co.column_dimensions[c].width=15
    header(co,6,['Primer fondeo','Cuentas requeridas','Ticket USD eq.','Origen externo','Ret30','Ret60','Ret90','Ret180','Ret365']+[f'M{i+1}'for i in range(12)])
    for j,c in enumerate(CCOLS):put(co,f'{c}7',f"='Supuestos'!{COLS[j]}31",'mmm-yy')
    for i,c in enumerate(COLS):
        row=9+i
        put(co,f'C{row}',f'=IF(COUNT(\'Supuestos\'!{c}31,\'Supuestos\'!$D$14,\'Supuestos\'!$D$36)=3,MAX(\'Supuestos\'!$D$36,DATE(YEAR(\'Supuestos\'!{c}31),MONTH(\'Supuestos\'!{c}31),\'Supuestos\'!$D$14)),"N/D")','dd/mm/yy')
        put(co,f'D{row}',f"='Requisitos'!{c}11",NUM);put(co,f'E{row}',"='Supuestos'!$D$8",MONEY);put(co,f'F{row}',f'=IF(COUNT(\'Supuestos\'!{c}44)=1,\'Supuestos\'!{c}44,"N/D")',PCT)
        for dest,src in zip('GHIJK',(37,38,39,40,41)):put(co,f'{dest}{row}',f'=IF(COUNT(\'Supuestos\'!$D${src})=1,\'Supuestos\'!$D${src},"N/D")',PCT)
        for j,cc in enumerate(CCOLS):
            age=f'{cc}{82+i}';ratio=f'{cc}{95+i}';put(co,age,f'=IF(COUNT({cc}$7,$C{row})=2,EOMONTH({cc}$7,0)-$C{row},"N/D")','0')
            curve=f'IF({age}<30,1+($G{row}-1)*{age}/30,IF({age}<60,$G{row}+($H{row}-$G{row})*({age}-30)/30,IF({age}<90,$H{row}+($I{row}-$H{row})*({age}-60)/30,IF({age}<180,$I{row}+($J{row}-$I{row})*({age}-90)/90,IF({age}<=365,$J{row}+($K{row}-$J{row})*({age}-180)/185,$K{row})))))'
            put(co,ratio,f'=IF(COUNT({age},$G{row}:$K{row},\'Supuestos\'!$D$45)=7,IF({j}>=\'Supuestos\'!$D$45,0,IF({age}<0,0,{curve})),"N/D")',PCT)
            put(co,f'{cc}{row}','='+guard([f'$D{row}',f'$E{row}',ratio],f'$D{row}*$E{row}*{ratio}'),MONEY)
            put(co,f'{cc}{30+i}','='+guard([f'{cc}{row}',f'$F{row}'],f'{cc}{row}*$F{row}'),MONEY)
            previous=f'{CCOLS[j-1]}{row}'if j else '0'
            formula=f'IF(COUNT($D{row},$E{row},$F{row},{cc}{row},{previous},\'Supuestos\'!$D$45)=6,IF({j}>=\'Supuestos\'!$D$45,0,IF($C{row}>EOMONTH({cc}$7,0),0,IF($C{row}>=DATE(YEAR({cc}$7),MONTH({cc}$7),1),($D{row}*$E{row}+{cc}{row})/2*(EOMONTH({cc}$7,0)-$C{row}+1)/DAY(EOMONTH({cc}$7,0)),({previous}+{cc}{row})/2))*$F{row}),"N/D")'
            put(co,f'{cc}{50+i}','='+formula,MONEY)
        for originrow in (30+i,50+i,82+i,95+i):put(co,f'C{originrow}',f'=C{row}','dd/mm/yy')
    for top,title in ((28,'Saldo externo al cierre por cohorte'),(48,'Saldo medio externo mensual por cohorte'),(80,'Edad en días al cierre'),(94,'Saldo relativo interpolado')):
        header(co,top,[title]+['']*8+[f'M{i+1}'for i in range(12)]);co.merge_cells(start_row=top,start_column=3,end_row=top,end_column=11);co[f'C{top}'].alignment=Alignment(horizontal='left',indent=1,vertical='center')
    for row,start,end,label in ((22,9,20,'Saldo total al cierre'),(43,30,41,'Saldo externo al cierre'),(63,50,61,'Saldo medio externo')):
        put(co,f'C{row}',label)
        for cc in CCOLS:put(co,f'{cc}{row}',f'=IF(COUNT({cc}{start}:{cc}{end})=12,SUM({cc}{start}:{cc}{end}),"N/D")',MONEY)
    header(co,66,['Primer fondeo','Cuentas','Saldo30d','Saldo60d','Saldo90d','Saldo180d','Saldo365d'])
    for i in range(12):
        origin=9+i;dest=67+i;put(co,f'C{dest}',f'=C{origin}','dd/mm/yy');put(co,f'D{dest}',f'=D{origin}',NUM)
        for cc,days,kn in (('E',30,'G'),('F',60,'H'),('G',90,'I'),('H',180,'J'),('I',365,'K')):put(co,f'{cc}{dest}',f'=IF(COUNT($C{origin},\'Supuestos\'!$D$45)=2,IF($C{origin}+{days}>EOMONTH(INDEX(\'Supuestos\'!$F$31:$Q$31,1,\'Supuestos\'!$D$45),0),"Fuera ventana",{guard([f"$D{origin}",f"$E{origin}",f"${kn}{origin}"],f"$D{origin}*$E{origin}*${kn}{origin}")}),"N/D")',MONEY)
    note(co,24,'USD equivalentes provisionales. Depósito bruto y saldo son stock, no beneficio. Finanzas valora el saldo medio externo dentro de 2027.',23)
    # Decision and financial compatibility. No fixed fee or fabricated unit economics.
    header(d,6,['Pregunta de inversión','Resultado activo','Unidad'])
    decisions={8:('Contribución contextual requerida',"='Supuestos'!D12",'cuentas',NUM),9:('Visitas requeridas bajo tasas de prueba',"=IF(COUNT('Requisitos'!F13:Q13)=12,SUM('Requisitos'!F13:Q13),\"N/D\")",'visitas',NUM),10:('Ratio requisito / panel parcial elegible',"=IF(COUNT('Requisitos'!F13:Q13,'Requisitos'!F16:Q16)=24,IF(SUM('Requisitos'!F16:Q16)>0,SUM('Requisitos'!F13:Q13)/SUM('Requisitos'!F16:Q16),\"N/D\"),\"N/D\")",'ratio panel',PCT),11:('Saldo medio externo en ventana activa',"=IF(COUNT('Requisitos'!F19:Q19,'Supuestos'!D45)=13,SUM('Requisitos'!F19:Q19)/'Supuestos'!D45,\"N/D\")",'USD eq.',MONEY),12:('Saldo externo al cierre de ventana',"=IF(COUNT('Supuestos'!D45)=1,INDEX('Requisitos'!F20:Q20,1,'Supuestos'!D45),\"N/D\")",'USD eq.',MONEY),14:('Valor neto banco según método/horizonte',"='Economia'!D8",'USD',MONEY),15:('Techo de inversión según retorno exigido',"='Economia'!D9",'USD',MONEY),16:('Piso operativo enfocado según costeo',"='Economia'!D15",'USD',MONEY),17:('Horas compatibles con el techo bancario',"='Economia'!D17",'horas',NUM),18:('ROI condicionado dentro de ventana',"='Requisitos'!Q26",'%',PCT),19:('Payback condicionado dentro de ventana',"=IF(COUNT('Requisitos'!F27:Q27)=12,IFERROR(MATCH(1,'Requisitos'!F27:Q27,0),\"No en ventana\"),\"N/D\")",'mes',NUM),20:('Fondeadas dentro de ventana, calidad',"=IF(COUNT('Requisitos'!F11:Q11)=12,SUM('Requisitos'!F11:Q11),\"N/D\")",'cuentas',NUM)}
    for row,(label,value,unit,fmt)in decisions.items():put(d,f'C{row}',label);put(d,f'D{row}',value,fmt);put(d,f'E{row}',unit).alignment=Alignment(indent=2,vertical='center');d.column_dimensions['E'].width=17
    for row in(24,):d.row_dimensions[row].height=45
    put(d,'C22','Fondeo potencial posterior (excluido)');put(d,'D22','=IF(COUNT(\'Requisitos\'!F12:Q12,\'Supuestos\'!D10,D20)=14,SUM(\'Requisitos\'!F12:Q12)*\'Supuestos\'!D10-D20,"N/D")',NUM)
    note(d,4,'La decisión parte de requisitos y valor bancario. Los escenarios previos 10/50/147 se retiran del uso con cliente; archivos históricos intactos.',11)
    note(d,21,'N/D conserva inputs faltantes. El techo es condicional a alcanzar la contribución; no es precio, presupuesto aprobado ni beneficio garantizado.',11)
    header(d,24,['Banda contextual','Cuentas según definición','Visitas necesarias','Saldo medio externo condicionado','Techo inversión condicionado'])
    for row,share in ((25,.01),(26,.05),(27,.10)):
        put(d,f'C{row}',share,PCT,input=True);put(d,f'D{row}',f'=C{row}*\'Supuestos\'!D6',NUM)
        put(d,f'E{row}',f'=IF(COUNT($D$8,$D$9)=2,IF($D$8>0,$D$9*D{row}/$D$8,"N/D"),"N/D")',NUM)
        put(d,f'F{row}',f'=IF(COUNT($D$8,$D$11)=2,IF($D$8>0,$D$11*D{row}/$D$8,"N/D"),"N/D")',MONEY)
        put(d,f'G{row}',f'=IF(COUNT($D$8,$D$15)=2,IF($D$8>0,$D$15*D{row}/$D$8,"N/D"),"N/D")',MONEY)
    note(d,29,'Comparación matemática bajo iguales tasas, calendario y origen; no son tres forecasts paralelos ni objetivos comerciales acordados.',11)
    note(d,31,'Programa enfocado candidato: 814h base +110h reserva=924h. Operations revisa carga, dependencias y disponibilidad; no se vende por horas.',11)
    note(d,32,'Si piso operativo supera techo bancario condicionado: revisar contribución/canales/alcance y evidencia; no inventar demanda o conversión.',11)
    note(d,34,'Una intervención de conversión puede mejorar tráfico de otros canales; requiere baseline y contrafactual propios para valorar incrementalidad.',11)
    note(d,35,'Educación/confianza y retención son mecanismos posibles, no ingresos adicionales sumables sin cohortes y deduplicación.',11)
    note(d,37,'M1 revisa accesos, medición y delivery. Las optimizaciones M3/6/9/12 sostienen el programa anual; no son cortes por ROI a tres meses.',11)
    note(d,39,'Riesgo central: tamaño/captura elegible y valor neto por saldo no están medidos. Cotización se prepara; la evidencia privada refina, no bloquea.',11)
    header(e,6,['Condición económica','Resultado','Unidad'])
    eco={8:('Valor banco si alcanza banda',"=IF(AND('Supuestos'!D16=\"Sí\",COUNT('Supuestos'!D17,'Decision'!D11)=2),'Decision'!D11*'Supuestos'!D17,\"N/D\")",MONEY),9:('Techo inversión para ROI exigido',"=IF(COUNT(D8,'Supuestos'!D18)=2,IF('Supuestos'!D18>=0,D8/(1+'Supuestos'!D18),\"N/D\"),\"N/D\")",MONEY),11:('Costo humano enfocado fully loaded',"=IF(AND('Supuestos'!D20=\"Sí\",COUNT('Supuestos'!D21,'Supuestos'!D23)=2,'Supuestos'!D21>0),'Supuestos'!D21*'Supuestos'!D23,\"N/D\")",MONEY),12:('Otros costos explícitos, sin doble overhead',"=IF(COUNT('Supuestos'!D25)=1,IF('Supuestos'!D25>=0,'Supuestos'!D25,\"N/D\"),\"N/D\")",MONEY),13:('Costo total interno condicionado','=IF(COUNT(D11:D12)=2,SUM(D11:D12),"N/D")',MONEY),15:('Piso comercial con margen autorizado',"=IF(AND('Supuestos'!D20=\"Sí\",COUNT(D13,'Supuestos'!D22)=2),'Supuestos'!D22,\"N/D\")",MONEY),17:('Horas compatibles con techo bancario',"=IF(COUNT(D9,D12,'Supuestos'!D21,'Supuestos'!D22)=4,IF(AND('Supuestos'!D20=\"Sí\",'Supuestos'!D21>0,'Supuestos'!D22>=0,'Supuestos'!D22<1),MAX(0,(D9*(1-'Supuestos'!D22)-D12)/'Supuestos'!D21),\"N/D\"),\"N/D\")",NUM),18:('Brecha techo menos piso','=IF(COUNT(D9,D15)=2,D9-D15,"N/D")',MONEY)}
    eco[15]=('Piso comercial con margen autorizado',"=IF(COUNT(D13,'Supuestos'!D22)=2,IF(AND('Supuestos'!D20=\"Sí\",'Supuestos'!D22>=0,'Supuestos'!D22<1),D13/(1-'Supuestos'!D22),\"N/D\"),\"N/D\")",MONEY)
    eco[11]=('Costo humano enfocado fully loaded',"=IF(AND('Supuestos'!D20=\"Sí\",COUNT('Supuestos'!D21,'Supuestos'!D23)=2,'Supuestos'!D21>=0),'Supuestos'!D21*'Supuestos'!D23,\"N/D\")",MONEY)
    eco[8]=('Valor banco por UN método/horizonte',"=IF('Supuestos'!D16=\"Sí\",IF('Supuestos'!D52=\"Saldo medio externo\",IF(COUNT('Supuestos'!D17,'Decision'!D11,'Supuestos'!D45)=3,'Decision'!D11*'Supuestos'!D17*'Supuestos'!D45/12,\"N/D\"),IF(COUNT('Supuestos'!D54,'Decision'!D20)=2,'Supuestos'!D54*'Decision'!D20,\"N/D\")),\"N/D\")",MONEY)
    for row,(label,value,fmt)in eco.items():put(e,f'C{row}',label);put(e,f'D{row}',value,fmt)
    note(e,20,'El piso cubre costos internos con margen autorizado; el techo responde al valor del banco. Ninguno es fee aprobado ni ahorro realizado.',11)
    header(e,23,['Tasa neta anual hipotética','Valor si banda se alcanza','Techo para retorno exigido','Naturaleza'])
    for row,rate in ((24,.01),(25,.02),(26,.04),(27,.06)):
        put(e,f'C{row}',rate,PCT,input=True);put(e,f'D{row}',f'=IF(COUNT(\'Decision\'!D11,\'Supuestos\'!D45)=2,\'Decision\'!D11*C{row}*\'Supuestos\'!D45/12,"N/D")',MONEY);put(e,f'E{row}',f'=IF(COUNT(D{row},\'Supuestos\'!D18)=2,IF(\'Supuestos\'!D18>=0,D{row}/(1+\'Supuestos\'!D18),"N/D"),"N/D")',MONEY);put(e,f'F{row}','Extremo; no recomendado'if rate==.06 else 'Experimento; no benchmark')
    note(e,29,'1/2/4/6% prueban aritmética del valor neto, no tasas bancarias observadas. 6% es extremo sin evidencia y no se recomienda.',11)
    header(e,32,['Capacidad','Base horas','Reserva horas','Total horas','Estado'])
    for row,vals in ((33,('Enfocado candidato',814,110,924,'Root: bottom-up, pendiente')), (34,('Foco previo histórico',1828,216,2044,'Cuestionado; no base actual')), (35,('Integral previo histórico',2560,336,2896,'Cuestionado; no base actual'))):
        for c,v in zip('CDEFG',vals):put(e,f'{c}{row}',v,NUM if c in'DEF'else None)
    note(e,37,'Mayor capacidad no produce por sí sola más cuentas. El candidato enfocado limita entregables y requiere implementación/QA coordinados.',11)
    note(e,39,'La valoración usa UN método Finance neto sobre saldo medio externo. No sumar FTP+ahorro si representan el mismo beneficio.',11)
    note(e,40,'Las transferencias internas son canibalización potencial; saldo bruto, depósito inicial y capital no son ingreso ni beneficio.',11)
    note(e,42,'Horizonte de valor 24/36m: sólo valor neto por fondeada aprobado por Finance (incluye descuento/costos). Sin tasa×2/3 ni cashflow futuro inventado.',11)
    note(e,43,'Servicio sigue12m. El método por cuenta no produce ROI/payback dentro del año; la distribución temporal de beneficios futuros está pendiente.',11)
    header(e,45,['Inversión hipotética, no quote','Valor mínimo por fondeada','Naturaleza'])
    for row,amount in ((46,10000),(47,25000),(48,50000)):
        put(e,f'C{row}',amount,MONEY,input=True);put(e,f'D{row}',f'=IF(COUNT(\'Decision\'!D20)=1,IF(\'Decision\'!D20>0,C{row}*(1+\'Supuestos\'!D18)/\'Decision\'!D20,"N/D"),"N/D")',MONEY);put(e,f'E{row}','Requisito; no valor observado')
    note(retro,5,'Jesús pidió qué aperturas habría podido cubrir SEO/GEO desde el lanzamiento de este año. Fecha exacta y ventana no están confirmadas.',11)
    note(retro,7,'Seleccionar Supuestos D29=Retrolanzamiento, cargar fecha D30 y meses D35. El mismo build calcula requisitos y maduración sólo en esa ventana.',11)
    note(retro,9,'Sin fecha/ventana, el modo retrospectivo propaga N/D. Programa2027 mantiene doce meses y no hereda una anualidad no declarada por el banco.',11)
    note(retro,11,'Las bandas152/760/1520 se interpretan como aperturas por defecto. Fondeo y saldo son calidad posterior; selector Fondeadas cambia el requisito.',11)
    note(retro,13,'La simulación es un contrafactual ilustrativo de requisitos; no aperturas históricas perdidas ni causalidad demostrada del lanzamiento.',11)
    note(retro,15,'No existe baseline histórico ni controles de exposición suficientes. Validar ventanas, condiciones de oferta, capacidad de indexación y funnel antes de atribuir.',11)
    note(retro,17,'No transferir un quote/valor aprobado de un contexto a otro sin validación. Los beneficios24/36m futuros se distinguen del saldo/valor dentro de ventana.',11)
    note(retro,19,'Para próximos lanzamientos: preparar arquitectura/contenidos/métricas antes del go-live, medir recorridos y optimizar cohortes durante el programa.',11)
    # Sources and all distinct channel mechanisms, with private baselines left absent.
    f.column_dimensions['C'].width=25;f.column_dimensions['D'].width=42;f.column_dimensions['E'].width=36;f.column_dimensions['F'].width=48;f.column_dimensions['G'].width=27
    put(f,'I5','Panel directo parcial');put(f,'J5',source['direct_panel_proxy_monthly'],NUM)
    header(f,7,['Carril','Evidencia disponible','Proxy/medida','Qué permitiría valorar','Baseline actual'])
    channels=[('Adquisición genérica','Panel curado de5 representantes',source['direct_panel_proxy_monthly'],'Captura elegible→visitas→fondeo; panel parcial','N/D'),('Adquisición condicional','Digital1300 + alto rendimientoPE390',1690,'Producto/intención/overlap antes de incluir','N/D'),('Educación/confianza','TREA/FSD/ahorro/transferencias',4070,'Asistencia medida; no sumarla a adquisición','N/D'),('Marca','140 histórico prelaunch: no baseline actual','Histórico separado','Protección/activación; visitas actuales necesarias','N/D'),('Conversión multicanal','Recorrido web→registro→app','N/D','Incremento respecto al contrafactual medido','N/D'),('Retención/origen externo','Cohortes/saldos banco pendientes','N/D','Valor neto incremental, sin doble conteo','N/D'),('Producto USD futuro','Consulta cuentaUSD',480,'No pertenece al producto actual público Soles','N/D')]
    for row,vals in enumerate(channels,8):
        for c,v in zip('CDEFG',vals):put(f,f'{c}{row}',v).alignment=Alignment(wrap_text=True,vertical='center')
        f.row_dimensions[row].height=38
    note(f,18,'3690 es una canasta de5 consultas de investigación; no censo del programa ni TAM. Dataset ampliado y familias nuevas se documentan aparte.',11)
    note(f,19,'Sin GSC/analytics, marca postlaunch, aperturas o cohortes: retornos adicionales N/D. No comprar APIs ni inflar proxies para cerrar un ROI.',11)
    header(f,22,['ID / naturaleza','Fuente','Alcance','Autoridad / estado'])
    entries=[('Meta provisional','Reunión06/10/2026 y operador','15200×2500≈US38m; periodo/definición pendientes','Banco no confirmó'),('Panel observado',source['source_file'],'PE/es, cinco representantes con overlap residual','DataForSEO; no TAM'),('Proyecciones retiradas','modelo-anual/results.json','10/50/147 no forecasts. Se conservan históricos','Withdrawn for client use'),('Capacidad candidata','Root: envelope bottom-up06/10','814h base +110h reserva, doce meses','Operations/Commercial pendientes'),('Hipótesis de requisitos','inputs.json','Tasas, distribución, saldos, origen y calendario','Pruebas; no benchmark'),('Valor bancario','Finance banco: método único pendiente','Tasa neta anual × saldo medio externo dentro2027','N/D'),('Costo/quote','Finance/Operations/Commercial pendientes','Fully loaded, otros costos, margen, fee, impuestos','N/D')]
    for row,vals in enumerate(entries,23):
        for c,v in zip('CDEF',vals):put(f,f'{c}{row}',v).alignment=Alignment(wrap_text=True,vertical='top')
        f.row_dimensions[row].height=45
    note(f,32,'Retirar de uso con cliente no elimina evidencia histórica: modelo-anual y modelo permanecen byte-for-byte intactos.',11)
    note(f,34,'M1 readiness/delivery; M3/6/9/12 optimización. Ejecución/caracterización técnica no garantizan rankings, citas ni apertura/fondeo.',11)
    wb.save(path)

def recalc(path,out):
    out.mkdir()
    with tempfile.TemporaryDirectory(prefix='pibank-decision-profile-')as profile:
        p=subprocess.run([SOFFICE,f'-env:UserInstallation={Path(profile).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=90)
        assert p.returncode==0 and(out/path.name).exists(),(p.stdout,p.stderr)
    return out/path.name

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--no-copy',action='store_true');args=parser.parse_args();cfg=json.loads((HERE/'inputs.json').read_text());source=json.loads((HERE/'sources.json').read_text())
    assert cfg['focused_capacity_hours_base']+cfg['focused_capacity_hours_reserve']==cfg['focused_capacity_hours_total']==924
    with tempfile.TemporaryDirectory(prefix='pibank-decision-build-')as tmp:
        raw=Path(tmp)/'raw';raw.mkdir();build(cfg,source,raw/NAME);final=recalc(raw/NAME,Path(tmp)/'calc');shutil.copy2(final,HERE/f'.{NAME}.building');(HERE/f'.{NAME}.building').replace(HERE/NAME)
    w=openpyxl.load_workbook(HERE/NAME,data_only=True);d=w['Decision'];e=w['Economia']
    results={'as_of':'2026-10-06','status':'Decision requirements; prior projections withdrawn for client use','forecast':False,'goal_status':'Contextual provisional bands, not agreed channel targets; period/definition of15200pending','selected_band_share':w['Supuestos']['D5'].value,'goal_definition':w['Supuestos']['G6'].value,'analysis_context':w['Supuestos']['D29'].value,'active_window_months':w['Supuestos']['D45'].value,'valuation_method':w['Supuestos']['D52'].value,'bank_value_horizon_months':w['Supuestos']['D53'].value,'selected_accounts_required':d['D8'].value,'funded_within_window_conditional':d['D20'].value,'visits_required_at_test_rates':d['D9'].value,'requirement_over_partial_panel_ratio':d['D10'].value,'average_external_balance_usd_equivalent_conditional':d['D11'].value,'external_balance_end_usd_equivalent_conditional':d['D12'].value,'bank_value_conditional_by_selected_method':d['D14'].value,'investment_ceiling_conditional':d['D15'].value,'operating_price_floor_conditional':d['D16'].value,'compatible_capacity_hours_conditional':d['D17'].value,'roi_conditional':d['D18'].value,'payback_conditional':d['D19'].value,
             'bands':[{'share':d[f'C{r}'].value,'accounts_required':d[f'D{r}'].value,'definition':w['Supuestos']['G6'].value,'visits_required_at_test_rates':d[f'E{r}'].value,'average_external_balance_conditional':d[f'F{r}'].value,'investment_ceiling_conditional':d[f'G{r}'].value}for r in(25,26,27)],'value_rate_experiments':[{'hypothetical_net_annual_rate':e[f'C{r}'].value,'conditional_value':e[f'D{r}'].value,'nature':e[f'F{r}'].value}for r in(24,25,26,27)],'focused_candidate_capacity':{'base_hours':814,'reserve_hours':110,'total_hours':924,'approved':False},'workbook':NAME,'valuation_basis':'ONE method: net annual rate×external average balance only within activewindow, OR Finance approved net present value per funded account over12/24/36m; no summation','required_net_value_per_funded_account':[{'hypothetical_investment_usd':e[f'C{r}'].value,'required_value_usd':e[f'D{r}'].value,'nature':'Arithmetic requirement, not quote or measured bank value'}for r in(46,47,48)]}
    results['funding_potential_after_window_excluded']=d['D22'].value
    for band in results['bands']:
        amount=d['D20'].value*band['accounts_required']/d['D8'].value if isinstance(d['D20'].value,(int,float))and d['D8'].value>0 else None
        band['funded_within_window_conditional']=amount
        band['required_net_value_per_funded_account']=[{'hypothetical_investment_usd':value,'required_value_usd':value*(1+w['Supuestos']['D18'].value)/amount if isinstance(amount,(int,float))and amount>0 else None,'nature':'Requirement, not quote or bank value'}for value in(10000,25000,50000)]
    (HERE/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2)+'\n')
    if not args.no_copy:
        OUT.mkdir(parents=True,exist_ok=True)
        for p in HERE.iterdir():
            if p.is_file()and p.suffix in('.json','.md','.py','.xlsx'):shutil.copy2(p,OUT/p.name)
    print(json.dumps(results,ensure_ascii=False))

if __name__=='__main__':main()
